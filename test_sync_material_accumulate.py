# -*- coding: utf-8 -*-
"""同一天同一物料出现在多张已确认单据时，同步应累加且重复同步幂等。"""
import os
import sys
import tempfile
import unittest

from openpyxl import Workbook, load_workbook

sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), "webapp"))
import app as webapp


DATE = "2026-08-27"


def _build_material(path):
    """最小物料明细表：肉类采购里有「牛肉」，数量列是该日，金额在数量列右侧两格。

    数量/金额预先写成 99/999，用来确认同步是用整天合计覆盖，而不是在旧值上再加。
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "肉类采购"
    serial = webapp._excel_serial(DATE)
    ws.cell(1, 1, "序号")
    ws.cell(1, 2, "物料代码")
    ws.cell(1, 4, "品名")
    ws.cell(1, 5, "单位")
    ws.cell(1, 6, serial)
    ws.cell(2, 6, "数量")
    ws.cell(2, 7, "单价")
    ws.cell(2, 8, "金额")
    ws.cell(3, 2, "N-BEEF")
    ws.cell(3, 4, "牛肉")
    ws.cell(3, 5, "斤")
    ws.cell(3, 6, 99)
    ws.cell(3, 8, 999)
    wb.save(path)


def _build_land(path):
    """同一天两张已确认单据：牛肉 10斤/500元、5斤/250元。"""
    wb = Workbook()
    ws = wb.active
    ws.title = "已落地"
    ws.append(["日期", "类别", "品名", "数量", "单价", "金额", "来源图片", "合计", "实付合计"])
    ws.append([DATE, "肉类", "牛肉", 10, 50, 500, "slip-a.jpg", 500, 500])
    ws.append([DATE, "肉类", "牛肉", 5, 50, 250, "slip-b.jpg", 250, 250])
    wb.save(path)


def _read_beef(path):
    wb = load_workbook(path)
    ws = wb["肉类采购"]
    return ws.cell(3, 6).value, ws.cell(3, 8).value


class SameDayMaterialSumTest(unittest.TestCase):
    def test_same_day_two_slips_sum_and_resync_is_idempotent(self):
        saved = dict(webapp._paths)
        try:
            with tempfile.TemporaryDirectory() as tmp:
                out = os.path.join(tmp, "out")
                os.makedirs(out)
                material = os.path.join(tmp, "物料明细.xlsx")
                _build_material(material)
                _build_land(os.path.join(out, webapp.LAND_FILENAME))
                webapp._paths["output_dir"] = out
                webapp._paths["material_file"] = material

                land = os.path.join(out, webapp.LAND_FILENAME)
                webapp.sync_to_material("unused-model")
                qty, amt = _read_beef(material)
                self.assertEqual((qty, amt), (15, 750))
                self.assertFalse(os.path.exists(land))

                # 同一批已确认单据再同步一次，不能变成 30/1500。
                # 成功同步会删掉中间表，所以用同一份落地表再写一次来验证覆盖而不是累加。
                _build_land(land)
                webapp.sync_to_material("unused-model")
                qty, amt = _read_beef(material)
                self.assertEqual((qty, amt), (15, 750))
                self.assertFalse(os.path.exists(land))
        finally:
            webapp._paths.clear()
            webapp._paths.update(saved)


if __name__ == "__main__":
    unittest.main()

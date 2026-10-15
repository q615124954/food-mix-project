# -*- coding: utf-8 -*-
"""肉类销货清单（文国副食，5 列粉纸复写）解析与阴影重试。

实拍三张（1726/1727/1728）复跑 OCR 的表格只有序号，品名/数量/单价/金额全空。
手写本身是「19.2×43=」写在数量列、单价留空，或数量/单价分开。这里用合成 HTML
锁住这两种能认出来的版式，并保证认不出字时复核页仍有可填的空行。不调用 OCR 接口。
"""
import os
import sys
import tempfile
import unittest

from PIL import Image, ImageDraw

ROOT = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.join(ROOT, "webapp"))
sys.path.insert(0, ROOT)

import app as webapp
import pipeline_min as core


def _table(rows):
    parts = ["<table><caption></caption>"]
    for row in rows:
        tds = "".join(f"<td>{c}</td>" for c in row)
        parts.append(f"<tr>{tds}</tr>")
    parts.append("</table>")
    return "".join(parts)


def _items(rows):
    return core.parse_line_items(core.parse_table(_table(rows)))


HEADER = ["序号", "规格及品名", "数量", "单价", "金额"]


def _skeleton(n, date_line="26年8月17日"):
    """复跑 OCR 的实际形态：表头 + n 行只有序号。"""
    body = "".join(
        f"<tr><td>{i}</td><td></td><td></td><td></td><td></td></tr>" for i in range(1, n + 1)
    )
    return (
        f"客户\n\n0005571\n\n{date_line}\n\n"
        "<table><caption></caption>"
        "<tr><td>序号</td><td>规格及品名</td><td>数量</td><td>单价</td><td>金额</td></tr>"
        f"{body}</table>\n\n览秀城"
    )


def _filled_meat():
    return _table([
        HEADER,
        ["1", "吊龙", "19.2×43=", "", "825"],
        ["2", "前夹", "20.5×34=", "", "697"],
        ["3", "牛骨", "", "28.7×2.5=", "71"],
        ["4", "牛油", "4.7", "5", "23"],
    ])


class MeatLayouts(unittest.TestCase):
    def test_compound_in_qty_column_empty_price(self):
        # 1726/1727/1728 的写法：19.2×43= 落在数量列，单价空着，金额在最后一列
        items = _items([
            HEADER,
            ["2", "吊龙", "19.2×43=", "", "825"],
            ["3", "前夹", "20.5x34=", "", "697"],
            ["4", "牛骨", "28.7X2.5=", "", "71"],
        ])
        self.assertEqual(
            [(it["品名"], it["数量"], it["金额"], it["类别"]) for it in items],
            [("吊龙", 19.2, 825, "肉类"), ("前夹", 20.5, 697, "肉类"), ("牛骨", 28.7, 71, "肉类")],
        )

    def test_compound_in_price_column(self):
        # 文档里的另一种：序号 | 品名 | 空 | 20×44 | 金额
        items = _items([HEADER, ["1", "吊龙", "", "20×44", "880"]])
        self.assertEqual(items[0]["数量"], 20)
        self.assertEqual(items[0]["金额"], 880)
        self.assertEqual(items[0]["类别"], "肉类")

    def test_separate_qty_and_price(self):
        items = _items([
            HEADER,
            ["1", "吊龙", "19.2", "43", "825"],
            ["2", "牛骨", "25.8", "2.5", "64"],
            ["3", "牛油", "4.7", "", "23"],  # 单价空，数量仍取数量列
        ])
        self.assertEqual([(it["品名"], it["数量"], it["金额"]) for it in items],
                         [("吊龙", 19.2, 825), ("牛骨", 25.8, 64), ("牛油", 4.7, 23)])

    def test_leading_times_is_not_compound_qty(self):
        # 单价列被认成「×25」时，不能把 25 当成数量
        items = _items([HEADER, ["1", "牛腩", "10", "×25", "250"]])
        self.assertEqual(items[0]["数量"], 10)
        self.assertEqual(items[0]["金额"], 250)

    def test_empty_name_keeps_qty_and_amount(self):
        items = _items([
            HEADER,
            ["3", "", "4.7×5=", "", "23"],
            ["4", "", "25.8", "", "64"],
        ])
        self.assertEqual(len(items), 2)
        self.assertEqual(items[0]["品名"], "")
        self.assertEqual(items[0]["数量"], 4.7)
        self.assertEqual(items[0]["金额"], 23)
        self.assertEqual(items[0]["类别"], "肉类")
        self.assertEqual(items[1]["数量"], 25.8)
        self.assertEqual(items[1]["金额"], 64)

    def test_numbered_empty_skeleton_becomes_placeholders(self):
        # 1726 复跑：5 行全空。1727/1728 复跑到 6 行。不能再解析成 0 行。
        for n in (5, 6):
            rows = core.parse_table(_skeleton(n))
            items = core.parse_line_items(rows)
            self.assertEqual(len(items), n, n)
            for it in items:
                self.assertEqual(it["品名"], "")
                self.assertIsNone(it["数量"])
                self.assertIsNone(it["金额"])
                self.assertEqual(it["类别"], "肉类")
            derived, stat = core.derive_price(items)
            self.assertEqual(stat["缺数量"], n)
            self.assertTrue(all(it["勾稽"] == "缺数量" for it in derived))

    def test_mixed_real_rows_and_blank_serials(self):
        items = _items([
            HEADER,
            ["1", "", "", "", ""],
            ["2", "吊龙", "19.2×43=", "", "825"],
            ["3", "", "", "", ""],
        ])
        self.assertEqual([it["品名"] for it in items], ["", "吊龙", ""])
        self.assertEqual(items[1]["数量"], 19.2)
        self.assertIsNone(items[0]["金额"])

    def test_header_total_and_unnumbered_amount_are_not_lines(self):
        items = _items([
            HEADER,
            ["合计", "", "", "", "1619"],
            ["", "", "", "", "1619"],
            ["2", "吊龙", "19.2", "43", "825"],
        ])
        self.assertEqual([it["品名"] for it in items], ["吊龙"])
        rows = core.parse_table(_table([
            HEADER,
            ["2", "吊龙", "19.2", "43", "825"],
            ["合计", "", "", "", "1619"],
        ]))
        self.assertEqual(core.extract_order_total(rows), 1619)

    def test_vegetable_seven_col_unchanged(self):
        rows = [
            ["编号", "商品名称", "单位", "数量", "单价", "金额", "备注"],
            ["1", "青菜", "斤", "2", "3", "6", ""],
            ["2", "", "斤", "1", "4", "4", ""],          # 品名空但有数量：保留
            ["3", "", "", "", "", "", ""],               # 只有序号：不进明细
            ["", "员工", "", "", "", "10", ""],
            ["4", "海底捞", "包", "1", "10", "10", ""],
        ]
        items = core.parse_line_items(rows)
        self.assertEqual([it["品名"] for it in items], ["青菜", "", "海底捞"])
        self.assertTrue(all(it["类别"] == "蔬菜" for it in items))
        self.assertFalse(items[0]["员工区"])
        self.assertFalse(items[1]["员工区"])
        self.assertTrue(items[2]["员工区"])
        self.assertEqual(items[1]["数量"], 1)
        self.assertEqual(items[1]["金额"], 4)
        self.assertEqual(core.count_staff_zone_missing(rows), 0)


class RicherParse(unittest.TestCase):
    def test_real_items_beat_skeleton_and_lost_table(self):
        skeleton = _items([HEADER] + [[str(i), "", "", "", ""] for i in range(1, 6)])
        filled = _items([HEADER, ["1", "吊龙", "19.2×43=", "", "825"]])
        self.assertTrue(core.ocr_parse_richer(filled, [HEADER], skeleton, [HEADER]))
        self.assertFalse(core.ocr_parse_richer(skeleton, [HEADER], filled, [HEADER]))
        # 1728 的增强把表格整段认丢：不能盖掉序号骨架
        self.assertFalse(core.ocr_parse_richer([], [], skeleton, [HEADER]))
        self.assertTrue(core.ocr_parse_richer(skeleton, [HEADER], [], []))

    def test_more_filled_fields_wins_ties_on_row_count(self):
        bare = [{"品名": "吊龙", "数量": None, "金额": None}]
        full = [{"品名": "吊龙", "数量": 19.2, "金额": 825}]
        self.assertTrue(core.ocr_parse_richer(full, [HEADER], bare, [HEADER]))
        self.assertFalse(core.ocr_parse_richer(bare, [HEADER], full, [HEADER]))

    def test_placeholders_do_not_count_as_content(self):
        blanks = [{"品名": "", "数量": None, "金额": None} for _ in range(5)]
        self.assertEqual(core.content_item_score(blanks), (0, 0))
        one = [{"品名": "", "数量": 4.7, "金额": 23}]
        self.assertEqual(core.content_item_score(one), (1, 2))


def _mean(img, box):
    crop = img.convert("L").crop(box)
    hist = crop.histogram()
    n = crop.size[0] * crop.size[1]
    return sum(i * hist[i] for i in range(256)) / n


def _stroke_contrast(img):
    """阴影里的笔画比它上方的纸面暗多少。"""
    return _mean(img, (30, 100, 180, 140)) - _mean(img, (30, 168, 180, 174))


def _pink_shadow(ink):
    img = Image.new("RGB", (360, 280), (236, 168, 190))
    px = img.load()
    for y in range(90, 280):
        for x in range(360):
            r, g, b = px[x, y]
            px[x, y] = (int(r * 0.28), int(g * 0.28), int(b * 0.28))
    ImageDraw.Draw(img).line([(16, 170), (340, 170)], fill=ink, width=4)
    return img


def _white_shadow_black():
    img = Image.new("RGB", (360, 280), (250, 250, 250))
    px = img.load()
    for y in range(90, 280):
        for x in range(360):
            r, g, b = px[x, y]
            px[x, y] = (int(r * 0.28), int(g * 0.28), int(b * 0.28))
    ImageDraw.Draw(img).line([(16, 170), (340, 170)], fill=(20, 20, 20), width=4)
    return img


class EnhancePasses(unittest.TestCase):
    def test_red_pass_lifts_blue_ink_on_pink_carbon(self):
        img = _pink_shadow((25, 35, 150))
        original = _stroke_contrast(img)
        red = core.enhance_image(img, "red")
        lifted = _stroke_contrast(red)
        self.assertLess(original, 25)
        self.assertGreater(lifted, 80)
        self.assertGreater(lifted, original * 5)
        # 手影（下半张）和亮处纸面被拉近，笔画才不会整块沉在黑里
        before = _mean(img, (40, 30, 160, 75)) / _mean(img, (40, 110, 160, 150))
        after_shadow = _mean(red, (40, 110, 160, 150))
        after_light = _mean(red, (40, 30, 160, 75))
        self.assertGreater(before, 2)
        self.assertGreater(after_shadow, 180)
        self.assertLess(after_light / after_shadow, 1.35)
        self.assertEqual(red.mode, "RGB")
        self.assertEqual(red.size, img.size)

    def test_flat_and_red_keep_black_ink_under_shadow(self):
        img = _white_shadow_black()
        original = _stroke_contrast(img)
        flat = _stroke_contrast(core.enhance_image(img, "flat"))
        red = _stroke_contrast(core.enhance_image(img, "red"))
        self.assertLess(original, 40)
        self.assertGreater(flat, 80)
        self.assertGreater(red, 80)

    def test_shadow_alias_and_prepare_image(self):
        img = _pink_shadow((25, 35, 150))
        self.assertEqual(
            core.enhance_image(img, True).tobytes(),
            core.enhance_image(img, "shadow").tobytes(),
        )
        with tempfile.TemporaryDirectory() as tmp:
            path = os.path.join(tmp, "ticket.jpg")
            img.save(path, "JPEG", quality=95)
            raw, name = core.prepare_image(path, enhance=False)
            shadow, _ = core.prepare_image(path, enhance=True)
            shadow_named, _ = core.prepare_image(path, enhance="shadow")
            red, _ = core.prepare_image(path, enhance="red")
        self.assertEqual(name, "ticket.jpg")
        self.assertEqual(shadow.getvalue(), shadow_named.getvalue())
        self.assertNotEqual(raw.getvalue(), red.getvalue())
        self.assertTrue(raw.getvalue().startswith(b"\xff\xd8"))


class OcrRetry(unittest.TestCase):
    def setUp(self):
        self._call = webapp.core.call_ocr

    def tearDown(self):
        webapp.core.call_ocr = self._call

    def test_skeleton_is_reviewable_and_later_pass_can_win(self):
        seen = []

        def fake(path, timeout=600, enhance=False):
            seen.append(enhance)
            if not enhance:
                return _skeleton(5)
            if enhance == "shadow":
                return "客户\n\n存根（白）\n"  # 1728：增强后表格没了
            if enhance == "red":
                return _filled_meat()  # 这一档没带日期
            raise AssertionError("red 已经认出明细，不应再打 flat")

        webapp.core.call_ocr = fake
        res = webapp.ocr_one("微信图片_20260902091935_1726_4.jpg")
        self.assertEqual(seen, [False, "shadow", "red"])
        self.assertEqual(res["date"], "2026-08-17")
        self.assertEqual([it["品名"] for it in res["items"]], ["吊龙", "前夹", "牛骨", "牛油"])
        self.assertEqual(res["items"][0]["数量"], 19.2)
        self.assertEqual(res["items"][0]["类别"], "肉类")
        self.assertFalse(any(s == "未识别到内容" for s in res["suspects"]))

    def test_all_passes_empty_still_leaves_placeholders(self):
        def fake(path, timeout=600, enhance=False):
            if enhance == "shadow":
                return "只有抬头没有表"
            return _skeleton(6, date_line="年8月20日")

        webapp.core.call_ocr = fake
        res = webapp.ocr_one("微信图片_20260902091937_1728_4.jpg")
        self.assertEqual(len(res["items"]), 6)
        self.assertTrue(all(it["品名"] == "" and it["数量"] is None for it in res["items"]))
        self.assertEqual(res["date"], "2026-08-20")
        self.assertTrue(any("疑似漏识别 6 行" in s for s in res["suspects"]))
        self.assertTrue(any("缺数量" in s for s in res["suspects"]))
        self.assertTrue(any("品名未识别" in s for s in res["suspects"]))
        self.assertFalse(any(s == "未识别到内容" for s in res["suspects"]))

    def test_shadow_failure_falls_through_to_flat(self):
        def fake(path, timeout=600, enhance=False):
            if not enhance or enhance == "red":
                return _skeleton(5)
            if enhance == "shadow":
                raise RuntimeError("暂时失败")
            if enhance == "flat":
                return _table([HEADER, ["1", "吊龙", "19.2×43=", "", "825"]])
            raise AssertionError(enhance)

        webapp.core.call_ocr = fake
        res = webapp.ocr_one("ticket.jpg")
        self.assertEqual([it["品名"] for it in res["items"]], ["吊龙"])
        self.assertEqual(res["date"], "2026-08-17")

    def test_partial_ticket_only_retries_shadow(self):
        seen = []
        original = _table([
            HEADER,
            ["1", "吊龙", "19.2×43=", "", "825"],
            ["2", "前夹", "20.5×34=", "", "697"],
            ["3", "", "", "", ""],
            ["4", "", "", "", ""],
        ])

        def fake(path, timeout=600, enhance=False):
            seen.append(enhance)
            if enhance == "shadow":
                return _table([
                    HEADER,
                    ["1", "吊龙", "19.2×43=", "", "825"],
                    ["2", "前夹", "20.5×34=", "", "697"],
                    ["3", "牛骨", "28.7×2.5=", "", "71"],
                    ["4", "牛油", "4.7", "5", "23"],
                ])
            if enhance in ("red", "flat"):
                raise AssertionError(enhance)
            return original

        webapp.core.call_ocr = fake
        res = webapp.ocr_one("ticket.jpg")
        self.assertEqual(seen, [False, "shadow"])
        self.assertEqual([it["品名"] for it in res["items"]], ["吊龙", "前夹", "牛骨", "牛油"])


class BlankPlaceholderLanding(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self._paths = dict(webapp._paths)
        webapp._paths["output_dir"] = self.tmp.name

    def tearDown(self):
        webapp._paths.clear()
        webapp._paths.update(self._paths)

    def test_unfilled_rows_stay_on_screen_but_do_not_land(self):
        from openpyxl import load_workbook
        items = [
            {"品名": "吊龙", "数量": 19.2, "单价": 42.97, "金额": 825, "类别": "肉类"},
            {"品名": "", "数量": 4.7, "单价": 4.89, "金额": 23, "类别": "肉类"},
            {"品名": "", "数量": None, "单价": None, "金额": None, "类别": "肉类"},
        ]
        webapp.append_land("1726.jpg", "2026-08-17", "肉类", 848, 848, items)
        wb = load_workbook(webapp.land_file())
        ws = wb.active
        names = [ws.cell(r, 3).value for r in range(2, ws.max_row + 1)]
        self.assertEqual(names, ["吊龙", None])
        self.assertEqual(ws.cell(3, 4).value, 4.7)
        self.assertEqual(ws.cell(3, 6).value, 23)


if __name__ == "__main__":
    unittest.main()

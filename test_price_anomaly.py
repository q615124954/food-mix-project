# -*- coding: utf-8 -*-
"""P0：数量丢小数点（鸡毛菜 5.8→58）要按单价异常标数量/金额存疑，不自动改数。

单价是 金额÷数量 反算的，58 斤 28 元反算出 0.48/斤，「数量×单价≠金额」永远不触发，
所以拿历史单价（物料表）或票面单价比，偏离超过 3 倍才标。
OCR 原文来自 samples/gold/ocr/（真实 HunyuanOCR 返回，可离线复现）。
"""
import os
import shutil
import sys
import tempfile
import unittest

BASE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.join(BASE, "webapp"))
import app as webapp
import pipeline_min as core

GOLD = os.path.join(BASE, "samples", "gold")


def _ocr_items(day):
    with open(os.path.join(GOLD, "ocr", f"receipt_2026-08-{day}.ocr.txt"), encoding="utf-8") as f:
        text = f.read()
    items = core.parse_line_items(core.parse_table(text))
    items, _ = core.derive_price(items)
    return items


class PriceAnomalyRule(unittest.TestCase):
    def test_hist_price_flags_lost_decimal(self):
        why = core.price_anomaly({"数量": 58, "金额": 28}, hist_price=4.83)
        self.assertIsNotNone(why)
        self.assertIn("小数点", why)

    def test_hist_price_ok_after_fix(self):
        self.assertIsNone(core.price_anomaly({"数量": 5.8, "金额": 28}, hist_price=4.83))

    def test_no_reference_no_flag(self):
        self.assertIsNone(core.price_anomaly({"数量": 58, "金额": 28}))

    def test_ticket_price_used_without_history(self):
        self.assertIsNotNone(core.price_anomaly({"数量": 58, "金额": 28, "票面单价": 4.8}))
        # 手写抹零（3×3.8=11.4 记 12）不算异常
        self.assertIsNone(core.price_anomaly({"数量": 3, "金额": 12, "票面单价": 3.8}))

    def test_history_overrides_bad_ticket_price(self):
        # 票面单价 3.9 被 OCR 成 39，但反算单价与历史一致 → 数量/金额没问题，不标
        self.assertIsNone(core.price_anomaly({"数量": 2, "金额": 8, "票面单价": 39}, hist_price=4.0))

    def test_ticket_column_reliability(self):
        bad_col = [{"数量": 2, "金额": 8, "票面单价": 39}, {"数量": 3, "金额": 18, "票面单价": 59},
                   {"数量": 1, "金额": 5, "票面单价": 48}, {"数量": 1, "金额": 10, "票面单价": 10}]
        self.assertFalse(core.ticket_price_reliable(bad_col))
        good_col = [{"数量": 58, "金额": 28, "票面单价": 4.8}, {"数量": 3, "金额": 12, "票面单价": 3.8},
                    {"数量": 1, "金额": 10, "票面单价": 10}, {"数量": 4, "金额": 27, "票面单价": 6.8}]
        self.assertTrue(core.ticket_price_reliable(good_col))


class GoldReceipt0813(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls._old = webapp._paths.get("material_file", "")
        webapp._paths["material_file"] = os.path.join(GOLD, "material_gold.xlsx")
        webapp.load_legal_names()

    @classmethod
    def tearDownClass(cls):
        webapp._paths["material_file"] = cls._old
        webapp.load_legal_names()

    def test_history_loaded(self):
        serial = webapp._excel_serial("2026-08-13")
        self.assertAlmostEqual(webapp._hist_prices["鸡毛菜"][serial], 28 / 5.8, places=2)
        # 当天那一格不算历史：金标表里鸡毛菜只有 8/13 有数，当天没有可用历史
        self.assertIsNone(webapp.hist_price("鸡毛菜", "2026-08-13"))
        self.assertAlmostEqual(webapp.hist_price("鸡毛菜"), 28 / 5.8, places=2)

    def test_jimaocai_flagged_only(self):
        items = _ocr_items("13")
        jm = [it for it in items if it["品名"] == "鸡毛菜"]
        self.assertEqual(len(jm), 1)
        self.assertEqual(jm[0]["数量"], 58)  # OCR 原值保留，不自动改
        suspects = webapp.compute_suspects("2026-08-13", items)
        self.assertTrue(jm[0]["数量存疑"])
        self.assertTrue(jm[0]["金额存疑"])
        self.assertIn("小数点", jm[0]["存疑说明"])
        self.assertTrue(any("单价异常" in s and "鸡毛菜" in s for s in suspects))
        # 其余有历史单价的物料行，数量/金额都不因单价被标（数量没回退）
        others = [it["品名"] for it in items
                  if it.get("存疑说明") and it["品名"] != "鸡毛菜"]
        self.assertEqual(others, [])

    def test_manual_fix_clears_flag(self):
        items = _ocr_items("13")
        jm = next(it for it in items if it["品名"] == "鸡毛菜")
        jm["数量"] = 5.8
        core.derive_price(items)
        webapp.compute_suspects("2026-08-13", items)
        self.assertFalse(jm["数量存疑"])
        self.assertAlmostEqual(jm["单价"], 4.83, places=2)

    def test_other_days_no_false_alarm(self):
        for day in ("11", "12"):
            items = _ocr_items(day)
            webapp.compute_suspects(f"2026-08-{day}", items)
            flagged = [it["品名"] for it in items
                       if it.get("存疑说明") and webapp.normalize_pin(it["品名"]) in webapp._legal_names]
            self.assertEqual(flagged, [], day)


class HistPriceExcludeAndOutlier(unittest.TestCase):
    def setUp(self):
        self._old = webapp._hist_prices
        s = webapp._excel_serial
        webapp._hist_prices = {"鸡毛菜": {
            s("2026-08-01"): 4.8, s("2026-08-02"): 5.0, s("2026-08-03"): 4.6,
            s("2026-08-05"): 0.48,          # 原表记错的离群值
            s("2026-08-13"): 28 / 58,       # 当天这一格（正是待核的错数）
        }}

    def tearDown(self):
        webapp._hist_prices = self._old

    def test_excludes_same_day_and_outlier(self):
        self.assertAlmostEqual(webapp.hist_price("鸡毛菜", "2026-08-13"), 4.8, places=2)

    def test_flag_and_clear(self):
        items = [{"品名": "鸡毛菜", "数量": 58, "金额": 28}]
        core.derive_price(items)
        webapp.compute_suspects("2026-08-13", items)
        self.assertTrue(items[0]["数量存疑"])
        self.assertIn("历史单价4.80", items[0]["存疑说明"])
        items[0]["数量"] = 5.8
        core.derive_price(items)
        webapp.compute_suspects("2026-08-13", items)
        self.assertFalse(items[0]["数量存疑"])


class UserOriginalSheet(unittest.TestCase):
    """用户原表：8/13 鸡毛菜误写成 58 斤 28 元。当天这一格不能把错数洗成「正常单价」。"""

    @classmethod
    def setUpClass(cls):
        from openpyxl import load_workbook
        cls.tmp = tempfile.mkdtemp()
        path = os.path.join(cls.tmp, "material_original.xlsx")
        wb = load_workbook(os.path.join(GOLD, "material_gold.xlsx"))
        ws = wb["日物料采购"]
        serial = webapp._excel_serial("2026-08-13")
        col = next(c for c in range(1, ws.max_column + 1)
                   if webapp._date_key(ws.cell(1, c).value) == serial)
        row = next(r for r in range(3, ws.max_row + 1) if ws.cell(r, 4).value == "鸡毛菜")
        ws.cell(row, col, 58)
        ws.cell(row, col + 2, 28)
        wb.save(path)
        cls._old = webapp._paths.get("material_file", "")
        webapp._paths["material_file"] = path
        webapp.load_legal_names()

    @classmethod
    def tearDownClass(cls):
        webapp._paths["material_file"] = cls._old
        webapp.load_legal_names()
        shutil.rmtree(cls.tmp, ignore_errors=True)

    def test_original_sheet_has_bad_value(self):
        serial = webapp._excel_serial("2026-08-13")
        self.assertAlmostEqual(webapp._hist_prices["鸡毛菜"][serial], 28 / 58, places=3)

    def test_ocr_58_flagged_then_fixed(self):
        items = _ocr_items("13")
        jm = next(it for it in items if it["品名"] == "鸡毛菜")
        webapp.compute_suspects("2026-08-13", items)
        self.assertTrue(jm["数量存疑"])
        self.assertTrue(jm["金额存疑"])
        jm["数量"] = 5.8
        core.derive_price(items)
        webapp.compute_suspects("2026-08-13", items)
        self.assertFalse(jm["数量存疑"])
        self.assertFalse(jm["金额存疑"])


if __name__ == "__main__":
    unittest.main()

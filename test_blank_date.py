# -*- coding: utf-8 -*-
"""日期为空的单据不得继承上一张单的日期，确认/同步都不得写入 Excel。"""
import copy
import os
import sys
import tempfile
import unittest

ROOT = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.join(ROOT, "webapp"))
sys.path.insert(0, ROOT)

import app as webapp
from fastapi import HTTPException
from openpyxl import Workbook, load_workbook

LAND_HEADERS = ["日期", "类别", "品名", "数量", "单价", "金额", "来源图片", "合计", "实付合计"]


def _merge_doc(ws, start, end):
    """与落地表一致：日期/类别/合计/实付合计按单据纵向合并，续行单元格为空。"""
    for col in (1, 2, 8, 9):
        ws.merge_cells(start_row=start, start_column=col, end_row=end, end_column=col)


class BlankDateTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self._paths = dict(webapp._paths)
        self._state = copy.deepcopy(webapp._state)
        self._state_file = webapp.STATE_FILE
        webapp.STATE_FILE = os.path.join(self.tmp.name, "state.json")
        webapp._paths["output_dir"] = os.path.join(self.tmp.name, "out")
        webapp._paths["material_file"] = os.path.join(self.tmp.name, "material.xlsx")
        os.makedirs(webapp._paths["output_dir"], exist_ok=True)

    def tearDown(self):
        webapp._paths.clear()
        webapp._paths.update(self._paths)
        webapp._state.clear()
        webapp._state.update(self._state)
        webapp.STATE_FILE = self._state_file

    def _write_material(self):
        """2026-08-01 / 2026-08-02 两列都预填数量，用来判断有没有被空日期单据盖掉。"""
        serial_801 = webapp._excel_serial("2026-08-01")
        serial_802 = webapp._excel_serial("2026-08-02")
        wb = Workbook()
        ws = wb.active
        ws.title = "日物料采购"
        ws.cell(1, 2, "代码")
        ws.cell(1, 4, "品名")
        ws.cell(1, 6, serial_801)
        ws.cell(1, 7, serial_802)
        # 数据从标题行 + 2 开始
        rows = [
            (3, "V1", "青菜", 99, 11),
            (4, "V2", "土豆", 99, 55),
            (5, "V3", "白菜", 77, 44),
        ]
        for r, code, name, qty_801, qty_802 in rows:
            ws.cell(r, 2, code)
            ws.cell(r, 4, name)
            ws.cell(r, 6, qty_801)
            ws.cell(r, 7, qty_802)
        wb.save(webapp._paths["material_file"])
        wb.close()

    def _write_land(self):
        """a.jpg 有日期（含合并续行）；b.jpg 整单日期为空；c.jpg 在空单之后，日期正常。"""
        wb = Workbook()
        ws = wb.active
        ws.title = "已落地"
        ws.append(LAND_HEADERS)
        ws.append(["2026-08-01", "蔬菜", "青菜", 3, 1, 3, "a.jpg", 7, 7])
        ws.append(["2026-08-01", "蔬菜", "土豆", 4, 1, 4, "a.jpg", 7, 7])
        ws.append(["", "蔬菜", "白菜", 8, 1, 8, "b.jpg", 15, 15])
        ws.append(["", "蔬菜", "青菜", 7, 1, 7, "b.jpg", 15, 15])
        ws.append(["2026-08-02", "蔬菜", "青菜", 6, 1, 6, "c.jpg", 11, 11])
        ws.append(["2026-08-02", "蔬菜", "土豆", 5, 1, 5, "c.jpg", 11, 11])
        _merge_doc(ws, 2, 3)
        _merge_doc(ws, 4, 5)
        _merge_doc(ws, 6, 7)
        path = webapp.land_file()
        wb.save(path)
        wb.close()
        return path

    def test_sync_blank_date_does_not_inherit_or_overwrite(self):
        self._write_material()
        self._write_land()

        results, _backup = webapp.sync_to_material("DeepSeek-V4-Pro")

        blank = [r for r in results if r["来源"] == "error" and "人工补填" in r["理由"]]
        self.assertEqual(sorted(r["落地品名"] for r in blank), ["白菜", "青菜"])
        for r in blank:
            self.assertEqual(r["日期"], "")
            self.assertIn("未写入", r["理由"])

        # 空日期单据的两行都没有写成成功记录
        written_blank = [r for r in results if r["来源"] != "error" and r["落地品名"] == "白菜"]
        self.assertEqual(written_blank, [])

        wb = load_workbook(webapp._paths["material_file"], data_only=True)
        ws = wb["日物料采购"]
        # 8月1日：a.jpg 正常写入；b.jpg 不得把白菜改成 8、也不得把青菜改成 7
        self.assertEqual(ws.cell(3, 6).value, 3)   # 青菜
        self.assertEqual(ws.cell(4, 6).value, 4)   # 土豆（同一张单的合并续行仍用本单日期）
        self.assertEqual(ws.cell(5, 6).value, 77)  # 白菜，原值保留
        # 8月2日：空单之后的 c.jpg 仍用自己的日期，续行不丢日期
        self.assertEqual(ws.cell(3, 7).value, 6)   # 青菜
        self.assertEqual(ws.cell(4, 7).value, 5)   # 土豆
        self.assertEqual(ws.cell(5, 7).value, 44)  # 白菜，原值保留
        wb.close()

    def test_confirm_blank_date_rejected_and_not_written(self):
        land = webapp.land_file()
        wb = Workbook()
        ws = wb.active
        ws.title = "已落地"
        ws.append(LAND_HEADERS)
        ws.append(["2026-08-01", "蔬菜", "原有", 1, 1, 1, "old.jpg", 1, 1])
        wb.save(land)
        wb.close()

        webapp._state["invoices"] = [{
            "filename": "b.jpg",
            "status": "pending",
            "ocr_done": True,
            "date": "",
            "items": [{"品名": "白菜", "数量": 8, "金额": 8}],
            "suspects": ["日期未识别"],
        }]
        payload = webapp.ConfirmPayload(
            date="  ",
            category="蔬菜",
            items=[{"品名": "白菜", "数量": 8, "金额": 8, "品名编辑": True, "原品名": "白采"}],
        )
        with self.assertRaises(HTTPException) as ctx:
            webapp.confirm("b.jpg", payload)
        self.assertEqual(ctx.exception.status_code, 400)
        self.assertIn("人工补填", ctx.exception.detail)
        self.assertIn("未写入", ctx.exception.detail)

        inv = webapp._state["invoices"][0]
        self.assertNotEqual(inv["status"], "confirmed")
        self.assertIn(webapp.BLANK_DATE_SUSPECT, inv["suspects"])
        self.assertIn("日期未识别", inv["suspects"])

        wb = load_workbook(land)
        ws = wb.active
        self.assertEqual(ws.max_row, 2)
        self.assertEqual(ws.cell(2, 3).value, "原有")
        self.assertEqual(ws.cell(2, 7).value, "old.jpg")
        wb.close()


if __name__ == "__main__":
    unittest.main()

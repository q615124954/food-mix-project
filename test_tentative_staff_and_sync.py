# -*- coding: utf-8 -*-
"""员工餐? 归并、财务中间表删除、实付可空。

不调在线模型。分层匹配要 LLM 时用假回复。
"""
import copy
import json
import os
import sys
import tempfile
import unittest

from openpyxl import Workbook, load_workbook
from openpyxl.workbook.workbook import Workbook as WorkbookClass

BASE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.join(BASE, "webapp"))
import app as webapp
import name_match
import pipeline_min as core

DATE = "2026-08-12"


def _material(path, serial):
    wb = Workbook()
    veg = wb.active
    veg.title = "日物料采购"
    veg.cell(1, 2, "代码")
    veg.cell(1, 4, "品名")
    veg.cell(1, 5, "单位")
    veg.cell(1, 6, serial)
    products = [
        ("S-GARLIC", "蒜头"),
        ("S-CLOVE", "去皮蒜米"),
        ("S-SOY", "生抽"),
        ("S-QING", "青菜"),
    ]
    for i, (code, name) in enumerate(products):
        row = 3 + i
        veg.cell(row, 2, code)
        veg.cell(row, 4, name)
        veg.cell(row, 5, "斤")
        veg.cell(row, 6, 99)
        veg.cell(row, 8, 999)
    staff = wb.create_sheet("员工餐")
    headers = ["序号", "日期", "类别", "金额", "序号", "日期", "类别", "金额", "总计"]
    for col, text in enumerate(headers, start=1):
        staff.cell(2, col, text)
    staff.cell(3, 2, serial)
    staff.cell(3, 4, 999)
    staff.cell(3, 6, serial)
    staff.cell(3, 8, 999)
    staff.cell(3, 9, 999)
    wb.save(path)
    wb.close()


def _land(path, rows):
    wb = Workbook()
    ws = wb.active
    ws.title = "已落地"
    ws.append(["日期", "类别", "品名", "数量", "单价", "金额", "来源图片",
               "合计", "实付合计", "归属", "品名编辑", "员工区"])
    for row in rows:
        ws.append(list(row))
    wb.save(path)
    wb.close()


def _cell(ws, name):
    for r in range(3, ws.max_row + 1):
        if ws.cell(r, 4).value == name:
            return ws.cell(r, 6).value, ws.cell(r, 8).value
    return None, None


class _Env(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self._paths = dict(webapp._paths)
        self._code_map = copy.deepcopy(webapp._code_map)
        self._name_dict = copy.deepcopy(webapp._name_dict)
        self._state = copy.deepcopy(webapp._state)
        self._variants = name_match._VARIANTS
        self._load_error = name_match._LOAD_ERROR
        self._variants_file = name_match.VARIANTS_FILE
        self._code_file = webapp.CODE_MAP_FILE
        self._dict_file = webapp.NAME_DICT_FILE
        self._save_state = webapp.save_state
        self._save_code = webapp.save_code_map
        self._chat = core.call_chat
        out = os.path.join(self.tmp.name, "out")
        os.makedirs(out)
        self.material = os.path.join(self.tmp.name, "物料明细.xlsx")
        self.serial = webapp._excel_serial(DATE)
        _material(self.material, self.serial)
        webapp._paths.update(output_dir=out, material_file=self.material)
        webapp._code_map = {"肉类": {}, "蔬菜": {}}
        webapp._name_dict = {"蔬菜": {}, "肉类": {}}
        webapp._state = {"invoices": []}
        webapp.CODE_MAP_FILE = os.path.join(self.tmp.name, "code_map.json")
        webapp.NAME_DICT_FILE = os.path.join(self.tmp.name, "product_dict.json")
        name_match.VARIANTS_FILE = os.path.join(self.tmp.name, "name_variants.json")
        name_match._VARIANTS = copy.deepcopy(name_match.DEFAULT_VARIANTS)
        name_match._LOAD_ERROR = ""
        webapp.save_state = lambda: None
        webapp.load_legal_names()
        self.llm_calls = []

        def chat(messages, model=None):
            self.llm_calls.append(messages)
            blob = json.dumps(messages, ensure_ascii=False)
            if "海天老抽" in blob or "消天老抽" in blob:
                return '{"目标品名":"生抽","置信度":0.99,"理由":"不该采纳"}'
            if "怪蒜低" in blob:
                return '{"目标品名":"蒜头","置信度":0.5,"理由":"拿不准"}'
            if "怪蒜边" in blob:
                return '{"目标品名":"蒜头","置信度":0.8,"理由":"刚好到门槛"}'
            if "怪蒜" in blob:
                return '{"目标品名":"蒜头","置信度":0.91,"理由":"就是蒜头"}'
            raise AssertionError("不应调用大模型")

        core.call_chat = chat
        self._real_layered = name_match.layered_match

        def layered(raw, candidates, **kwargs):
            forced = {
                "怪蒜": (["蒜头", "青菜"], 0.85),
                "怪蒜低": (["蒜头", "青菜"], 0.85),
                "怪蒜边": (["蒜头", "青菜"], 0.8),
                "海天老抽": (["生抽", "青菜"], 0.99),
            }
            if raw in forced:
                cands, conf = forced[raw]
                return {"action": "llm", "name": raw, "suggestion": cands[0], "source": "llm",
                        "confidence": conf, "llm_candidates": list(cands)}
            return self._real_layered(raw, candidates, **kwargs)

        name_match.layered_match = layered

    def tearDown(self):
        core.call_chat = self._chat
        name_match.layered_match = self._real_layered
        name_match.VARIANTS_FILE = self._variants_file
        name_match._VARIANTS = self._variants
        name_match._LOAD_ERROR = self._load_error
        webapp.save_state = self._save_state
        webapp.save_code_map = self._save_code
        webapp.CODE_MAP_FILE = self._code_file
        webapp.NAME_DICT_FILE = self._dict_file
        webapp._code_map = self._code_map
        webapp._name_dict = self._name_dict
        webapp._state = self._state
        webapp._paths.clear()
        webapp._paths.update(self._paths)
        webapp.load_legal_names()

    def _sync(self, rows, paid=None):
        total = sum(row[2] for row in rows)
        paid_v = total if paid is None else paid
        land_rows = []
        for name, qty, amt, attr, edited, zone in rows:
            land_rows.append([DATE, "蔬菜", name, qty, None, amt, "slip.jpg",
                              total, paid_v, attr, edited, zone])
        _land(webapp.land_file(), land_rows)
        return webapp.sync_to_material("unused-model")

    def _read(self):
        wb = load_workbook(self.material, data_only=True)
        veg = wb["日物料采购"]
        staff = wb["员工餐"]
        got = {
            "蒜头": _cell(veg, "蒜头"),
            "去皮蒜米": _cell(veg, "去皮蒜米"),
            "生抽": _cell(veg, "生抽"),
            "青菜": _cell(veg, "青菜"),
            "员工餐": staff.cell(3, 4).value,
            "日物料": staff.cell(3, 8).value,
        }
        wb.close()
        return got


class TentativeStaffMatch(_Env):
    def test_high_confidence_variant_lands_in_daily_material(self):
        """大蒜头/蒜米 是员工餐?，分层别名高于 0.8，进日物料，不调模型。"""
        self._sync([
            ("大蒜头", 3, 18, "员工餐?", False, False),
            ("蒜米", 1, 5, "员工餐?", False, False),
            ("野草", 1, 7, "员工餐?", False, False),
        ])
        got = self._read()
        self.assertEqual(got["蒜头"], (3, 18))
        self.assertEqual(got["去皮蒜米"], (1, 5))
        self.assertEqual(got["日物料"], 23)
        self.assertEqual(got["员工餐"], 7)
        self.assertEqual(self.llm_calls, [])
        garlic = name_match.variants_for("蔬菜").get("蒜头") or []
        self.assertIn("大蒜头", garlic)

    def test_llm_above_threshold_lands_and_below_stays_staff(self):
        self._sync([
            ("怪蒜", 2, 12, "员工餐?", False, False),
            ("怪蒜低", 1, 4, "员工餐?", False, False),
        ])
        got = self._read()
        self.assertEqual(got["蒜头"], (2, 12))
        self.assertEqual(got["员工餐"], 4)
        self.assertEqual(got["日物料"], 12)
        self.assertEqual(len(self.llm_calls), 2)
        self.assertIn("怪蒜", name_match.variants_for("蔬菜").get("蒜头") or [])
        self.assertNotIn("怪蒜低", name_match.variants_for("蔬菜").get("蒜头") or [])

    def test_existing_sync_threshold_accepts_point_eight(self):
        """与同步循环相同：置信度 0.8 落地，低于 0.8 不落地。"""
        self.assertEqual(webapp.SYNC_CONF, 0.8)
        self._sync([("怪蒜边", 1, 6, "员工餐?", False, False)])
        self.assertEqual(self._read()["蒜头"], (1, 6))
        self.assertEqual(self._read()["员工餐"], 0)

    def test_haidilao_and_staff_zone_never_call_llm(self):
        calls = []
        real = name_match.layered_match

        def spy(raw, candidates, **kwargs):
            calls.append(raw)
            return real(raw, candidates, **kwargs)

        name_match.layered_match = spy
        try:
            self._sync([
                ("海底捞", 1, 40, "员工餐", False, False),
                ("五花肉", 1, 21, "员工餐?", False, True),
                ("青菜", 2, 10, "物料", False, False),
            ])
        finally:
            name_match.layered_match = layered_restore(self)
        got = self._read()
        self.assertEqual(got["青菜"], (2, 10))
        self.assertEqual(got["日物料"], 10)
        self.assertEqual(got["员工餐"], 61)
        self.assertNotIn("海底捞", calls)
        self.assertNotIn("五花肉", calls)
        self.assertEqual(self.llm_calls, [])

    def test_dark_soy_blocked_even_when_model_says_high_confidence(self):
        hit = webapp.match_tentative_to_catalog(
            "海天老抽", "蔬菜", ["生抽", "青菜"], {"生抽": "S-SOY", "青菜": "S-QING"}, "m")
        self.assertIsNone(hit)
        self.assertTrue(self.llm_calls)  # 候选里还有青菜，模型被问过，但生抽答案作废
        self.llm_calls.clear()

        # 候选只剩生抽时，连模型都不问
        def only_soy(raw, candidates, **kwargs):
            return {"action": "llm", "name": raw, "suggestion": "生抽", "source": "llm",
                    "confidence": 0.99, "llm_candidates": ["生抽"]}

        name_match.layered_match = only_soy
        try:
            hit = webapp.match_tentative_to_catalog(
                "海天老抽", "蔬菜", ["生抽"], {"生抽": "S-SOY"}, "m")
            detail = webapp.describe_tentative_match(
                "海天老抽", "蔬菜", ["生抽"], {"生抽": "S-SOY"}, "m")
        finally:
            name_match.layered_match = layered_restore(self)
        self.assertIsNone(hit)
        self.assertEqual(self.llm_calls, [])
        self.assertFalse(detail["land"])
        self.assertAlmostEqual(detail["置信度"], 0.99)
        self.assertIn("未调用大模型", detail["理由"])
        self.assertIn("生抽", detail["理由"])
        self.assertNotIn("低于", detail["理由"])

        self._sync([("海天老抽", 1, 13, "员工餐?", False, False)])
        got = self._read()
        self.assertEqual(got["生抽"], (99, 999))
        self.assertEqual(got["员工餐"], 13)
        self.assertEqual(got["日物料"], 0)
        self.assertFalse(name_match.add_variant("蔬菜", "生抽", "海天老抽", code="S-SOY"))

    def test_human_edited_name_is_not_overwritten(self):
        self._sync([("大蒜头", 3, 18, "员工餐?", True, False)])
        got = self._read()
        self.assertEqual(got["蒜头"], (99, 999))
        self.assertEqual(got["员工餐"], 18)
        self.assertEqual(self.llm_calls, [])

        webapp._state = {"invoices": [{
            "filename": "edit.jpg", "status": "pending", "ocr_done": True, "items": [],
        }]}
        items = [{"品名": "大蒜头", "原品名": "大蒜头", "品名编辑": True,
                  "数量": 3, "金额": 18, "归属待确认": True}]
        payload = webapp.ConfirmPayload(date=DATE, category="蔬菜", items=items, paid_total=18)
        self.assertEqual(webapp.confirm("edit.jpg", payload), {"ok": True})
        saved = webapp._state["invoices"][0]["items"][0]
        self.assertEqual(saved["品名"], "大蒜头")
        ws = load_workbook(webapp.land_file()).active
        self.assertEqual(ws.cell(2, 3).value, "大蒜头")
        self.assertEqual(ws.cell(2, 10).value, "员工餐")

    def _add_product(self, code, name):
        wb = load_workbook(self.material)
        veg = wb["日物料采购"]
        row = veg.max_row + 1
        veg.cell(row, 2, code)
        veg.cell(row, 4, name)
        veg.cell(row, 5, "斤")
        veg.cell(row, 6, 99)
        veg.cell(row, 8, 999)
        wb.save(self.material)
        wb.close()

    def test_mid_score_suggestion_keeps_real_confidence(self):
        """腱子肉对猪腱子肉约 0.86：只建议、不落地，结果表记这个分数，而不是 0 / 低于 0.8。"""
        self._add_product("M-PORK", "猪腱子肉")
        catalog = ["蒜头", "去皮蒜米", "生抽", "青菜", "猪腱子肉"]
        expected = name_match.layered_match("腱子肉", catalog, category="蔬菜")
        self.assertEqual(expected["action"], "suggest")
        self.assertGreater(expected["confidence"], 0.84)
        self.assertLess(expected["confidence"], 0.88)
        results, _backup = self._sync([("腱子肉", 1, 16, "员工餐?", False, False)])
        row = _product_result(results, "腱子肉")
        self.assertAlmostEqual(row["置信度"], expected["confidence"])
        self.assertEqual(row["目标品名"], "员工餐")
        self.assertIn("猪腱子肉", row["理由"])
        self.assertIn("建议", row["理由"])
        self.assertIn("未调用大模型", row["理由"])
        self.assertNotIn("低于", row["理由"])
        self.assertNotIn("未达到同步置信度", row["理由"])
        self.assertEqual(self.llm_calls, [])
        got = self._read()
        self.assertEqual(got["员工餐"], 16)
        self.assertEqual(got["日物料"], 0)
        wb = load_workbook(self.material, data_only=True)
        self.assertEqual(_cell(wb["日物料采购"], "猪腱子肉"), (99, 999))
        wb.close()

    def test_success_records_score_that_caused_landing(self):
        """别名落地记 0.95，模型落地记模型分；本来就精确的品名才记 1。"""
        results, _backup = self._sync([
            ("大蒜头", 3, 18, "员工餐?", False, False),
            ("怪蒜", 2, 12, "员工餐?", False, False),
            ("青菜", 1, 8, "物料", False, False),
        ])
        garlic = _product_result(results, "大蒜头")
        self.assertEqual(garlic["目标品名"], "蒜头")
        self.assertAlmostEqual(garlic["置信度"], 0.95)
        self.assertNotEqual(garlic["置信度"], 1)
        self.assertIn("别名", garlic["理由"])
        model_hit = _product_result(results, "怪蒜")
        self.assertEqual(model_hit["目标品名"], "蒜头")
        self.assertAlmostEqual(model_hit["置信度"], 0.91)
        self.assertNotEqual(model_hit["置信度"], 1)
        exact = _product_result(results, "青菜")
        self.assertEqual(exact["目标品名"], "青菜")
        self.assertEqual(exact["置信度"], 1)
        self.assertIn("精确", exact["理由"])

    def test_llm_below_threshold_records_model_score(self):
        results, _backup = self._sync([("怪蒜低", 1, 4, "员工餐?", False, False)])
        row = _product_result(results, "怪蒜低")
        self.assertAlmostEqual(row["置信度"], 0.5)
        self.assertEqual(row["目标品名"], "员工餐")
        self.assertIn("大模型", row["理由"])
        self.assertIn("0.5", row["理由"])
        self.assertIn("0.8", row["理由"])
        self.assertNotIn("未达到同步置信度", row["理由"])
        self.assertEqual(self._read()["员工餐"], 4)
        self.assertEqual(self._read()["蒜头"], (99, 999))

    def test_no_close_candidate_is_not_blanket_below_threshold(self):
        results, _backup = self._sync([("野草", 1, 7, "员工餐?", False, False)])
        row = _product_result(results, "野草")
        expected = name_match.layered_match(
            "野草", ["蒜头", "去皮蒜米", "生抽", "青菜"], category="蔬菜")
        self.assertEqual(expected["action"], "keep")
        self.assertAlmostEqual(row["置信度"], expected["confidence"])
        self.assertIn("没有接近的候选", row["理由"])
        self.assertIn("未调用大模型", row["理由"])
        self.assertNotIn("低于", row["理由"])
        self.assertEqual(self.llm_calls, [])

    def test_blocked_soy_keeps_rejected_score(self):
        results, _backup = self._sync([("海天老抽", 1, 13, "员工餐?", False, False)])
        row = _product_result(results, "海天老抽")
        self.assertAlmostEqual(row["置信度"], 0.99)
        self.assertEqual(row["目标品名"], "员工餐")
        self.assertIn("生抽", row["理由"])
        self.assertIn("未采纳", row["理由"])
        self.assertNotIn("低于", row["理由"])
        self.assertEqual(self._read()["生抽"], (99, 999))
        self.assertEqual(self._read()["员工餐"], 13)

    def test_skipped_match_confidence_stays_blank(self):
        """真员工餐、员工区、人工改名不跑匹配，置信度留空，不填 0 或 1。"""
        results, _backup = self._sync([
            ("海底捞", 1, 40, "员工餐", False, False),
            ("五花肉", 1, 21, "员工餐?", False, True),
            ("大蒜头", 3, 18, "员工餐?", True, False),
            ("青菜", 2, 10, "员工餐", False, False),
        ])
        haidilao = _product_result(results, "海底捞")
        zone = _product_result(results, "五花肉")
        edited = _product_result(results, "大蒜头")
        confirmed = _product_result(results, "青菜")
        self.assertIsNone(haidilao["置信度"])
        self.assertIn("真员工餐", haidilao["理由"])
        self.assertIn("海底捞", haidilao["理由"])
        self.assertIsNone(zone["置信度"])
        self.assertIn("员工区", zone["理由"])
        self.assertIsNone(edited["置信度"])
        self.assertIn("人工", edited["理由"])
        self.assertIsNone(confirmed["置信度"])
        self.assertIn("真员工餐", confirmed["理由"])
        self.assertEqual(self._read()["蒜头"], (99, 999))
        self.assertEqual(self._read()["青菜"], (99, 999))
        self.assertEqual(self.llm_calls, [])


def _product_result(results, name):
    matched = [row for row in results if row.get("落地品名") == name]
    if not matched:
        raise AssertionError(name)
    return matched[0]


def layered_restore(env):
    """setUp 里包过 layered_match。海带捞用例中途拨开后要装回去。"""
    real = env._real_layered

    def layered(raw, candidates, **kwargs):
        forced = {
            "怪蒜": (["蒜头", "青菜"], 0.85),
            "怪蒜低": (["蒜头", "青菜"], 0.85),
            "怪蒜边": (["蒜头", "青菜"], 0.8),
            "海天老抽": (["生抽", "青菜"], 0.99),
        }
        if raw in forced:
            cands, conf = forced[raw]
            return {"action": "llm", "name": raw, "suggestion": cands[0], "source": "llm",
                    "confidence": conf, "llm_candidates": list(cands)}
        return real(raw, candidates, **kwargs)

    return layered


class PaidAndIntermediateFile(_Env):
    def test_empty_paid_uses_sum_and_flags_hidden_discount(self):
        items = [
            {"品名": "青菜", "数量": 2, "金额": 100},
            {"品名": "海底捞", "数量": 1, "金额": 40, "员工区": False},
        ]
        webapp._state = {"invoices": [{
            "filename": "pay.jpg", "status": "pending", "ocr_done": True,
            "items": items, "order_total": 100,
        }]}
        payload = webapp.ConfirmPayload(date=DATE, category="蔬菜", items=items)
        self.assertEqual(webapp.confirm("pay.jpg", payload), {"ok": True})
        inv = webapp._state["invoices"][0]
        self.assertEqual(inv["paid_total"], 140)
        self.assertEqual(inv["total"], 140)
        self.assertTrue(any("优惠" in s and "员工餐" in s for s in inv["suspects"]))
        ws = load_workbook(webapp.land_file()).active
        self.assertEqual(ws.cell(2, 9).value, 140)

        # 按合计入账时员工餐仍是海底捞的 40，不会被票面更低的合计抹成 0
        webapp.sync_to_material("unused-model")
        got = self._read()
        self.assertEqual(got["青菜"], (2, 100))
        self.assertEqual(got["员工餐"], 40)
        self.assertEqual(got["日物料"], 100)

        # 用户亲手把实付改成优惠后的数：不再标这条存疑，员工餐按实付减日物料
        webapp._state = {"invoices": [{
            "filename": "pay2.jpg", "status": "pending", "ocr_done": True,
            "items": items, "order_total": 100,
        }]}
        payload = webapp.ConfirmPayload(
            date=DATE, category="蔬菜", items=items, paid_total=100)
        self.assertEqual(webapp.confirm("pay2.jpg", payload), {"ok": True})
        inv = webapp._state["invoices"][0]
        self.assertEqual(inv["paid_total"], 100)
        self.assertEqual(inv["suspects"], [])
        webapp.sync_to_material("unused-model")
        got = self._read()
        self.assertEqual(got["员工餐"], 0)
        self.assertEqual(got["日物料"], 100)

    def test_empty_paid_without_lower_ticket_total_is_not_suspect(self):
        items = [{"品名": "青菜", "数量": 1, "金额": 10}]
        webapp._state = {"invoices": [{
            "filename": "plain.jpg", "status": "pending", "ocr_done": True, "items": items,
        }]}
        payload = webapp.ConfirmPayload(date=DATE, category="蔬菜", items=items)
        self.assertEqual(webapp.confirm("plain.jpg", payload), {"ok": True})
        inv = webapp._state["invoices"][0]
        self.assertEqual(inv["paid_total"], 10)
        self.assertEqual(inv["suspects"], [])

    def test_intermediate_excel_deleted_only_after_successful_sync(self):
        rows = [("青菜", 2, 10, "物料", False, False)]
        self._sync(rows)
        self.assertFalse(os.path.exists(webapp.land_file()))
        self.assertEqual(self._read()["青菜"], (2, 10))

        # 校验失败：中间表还在，物料表回到这次写入之前（不能留下 4/20）
        land = webapp.land_file()
        _land(land, [[DATE, "蔬菜", "青菜", 4, None, 20, "slip.jpg", 20, 20, "物料", False, False]])
        real_validate = webapp._validate_material_write

        def fail_validate(*_a, **_k):
            raise ValueError("读回不一致")

        webapp._validate_material_write = fail_validate
        try:
            with self.assertRaises(webapp.HTTPException) as ctx:
                webapp.sync_to_material("unused-model")
        finally:
            webapp._validate_material_write = real_validate
        self.assertEqual(ctx.exception.status_code, 500)
        self.assertIn("财务中间表已保留", ctx.exception.detail)
        self.assertIn("重试", ctx.exception.detail)
        self.assertTrue(os.path.exists(land))
        self.assertEqual(self._read()["青菜"], (2, 10))

        # 删除失败：物料数保留这次写入的 4/20，接口报可重试的错，中间表还在
        _land(land, [[DATE, "蔬菜", "青菜", 4, None, 20, "slip.jpg", 20, 20, "物料", False, False]])
        real_remove = os.remove

        def fail_remove(path):
            if os.path.abspath(path) == os.path.abspath(land):
                raise PermissionError("occupied")
            return real_remove(path)

        os.remove = fail_remove
        try:
            with self.assertRaises(webapp.HTTPException) as ctx:
                webapp.sync_to_material("unused-model")
        finally:
            os.remove = real_remove
        self.assertEqual(ctx.exception.status_code, 500)
        self.assertIn("删除失败", ctx.exception.detail)
        self.assertIn("重试", ctx.exception.detail)
        self.assertTrue(os.path.exists(land))
        self.assertEqual(self._read()["青菜"], (4, 20))

    def test_sync_save_failure_keeps_intermediate_file(self):
        land = webapp.land_file()
        _land(land, [[DATE, "蔬菜", "青菜", 1, None, 3, "slip.jpg", 3, 3, "物料", False, False]])
        real_save = WorkbookClass.save

        def locked(self, path):
            if str(path) == self and False:
                pass
            raise PermissionError("locked")

        WorkbookClass.save = locked
        try:
            with self.assertRaises(PermissionError):
                webapp.sync_to_material("unused-model")
        finally:
            WorkbookClass.save = real_save
        self.assertTrue(os.path.exists(land))
        self.assertEqual(self._read()["青菜"], (99, 999))


if __name__ == "__main__":
    unittest.main()

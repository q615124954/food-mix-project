# -*- coding: utf-8 -*-
"""美菜订单截图：实收解析、按日合并、长图分条。不调在线 OCR。

美菜复核明细：品名和单位都相同的行加总数量和金额（单张图一次，按日合并后再来一次）。
手写单不按品名合并。

金标 samples/gold/meicai_2026-09-09.json 是复核落地后的口径，不是 OCR 像素猜测。
OCR 文本里没有「收」时，数量和金额留空并标存疑。
Shot1 / Shot2 金标同样是复核后的落地值。
samples/gold/ocr/meicai_shot1.ocr.txt、meicai_shot2.ocr.txt 是按本机错认形态整理的离线文本
（香菜实收被下一件货接走、白壳中码、汽水碳酸饮料、促重行、美菜鲜坊、空数量小葱），
解析结果必须和金标的行数、品名、数量、金额一致。
samples/gold/ocr/meicai_shot{1,2,3}.live.ocr.txt 是本机重跑的 OCR 原文，门禁以这三份为准。
Shot2 的弱 OCR（没有「收」行）必须标存疑，不能把金标里的实收数编出来。
"""
import copy
import json
import os
import re
import sys
import tempfile
import unittest

from PIL import Image

ROOT = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.join(ROOT, "webapp"))
sys.path.insert(0, ROOT)

import pipeline_min as core
import app as webapp
from fastapi import HTTPException
from openpyxl import load_workbook

GOLD_JSON = os.path.join(ROOT, "samples", "gold", "meicai_2026-09-09.json")
OCR_TXT = os.path.join(ROOT, "samples", "gold", "ocr", "meicai_2026-09-09.ocr.txt")
HAND_TXT = os.path.join(ROOT, "samples", "gold", "ocr", "receipt_2026-08-11.ocr.txt")
SHOT_GOLD = {
    "shot1": os.path.join(ROOT, "samples", "gold", "meicai_shot1.json"),
    "shot2": os.path.join(ROOT, "samples", "gold", "meicai_shot2.json"),
}
SHOT_OCR = {
    "shot1": os.path.join(ROOT, "samples", "gold", "ocr", "meicai_shot1.ocr.txt"),
    "shot2": os.path.join(ROOT, "samples", "gold", "ocr", "meicai_shot2.ocr.txt"),
}
SHOT3_OCR = os.path.join(ROOT, "samples", "gold", "ocr", "meicai_shot3.ocr.txt")
SHOT3_TILED_OCR = os.path.join(ROOT, "samples", "gold", "ocr", "meicai_shot3.tiled.ocr.txt")
LIVE_OCR = {
    "shot1": os.path.join(ROOT, "samples", "gold", "ocr", "meicai_shot1.live.ocr.txt"),
    "shot2": os.path.join(ROOT, "samples", "gold", "ocr", "meicai_shot2.live.ocr.txt"),
    "shot3": os.path.join(ROOT, "samples", "gold", "ocr", "meicai_shot3.live.ocr.txt"),
}


def _ocr_text():
    with open(OCR_TXT, encoding="utf-8") as f:
        return f.read()


def _gold():
    with open(GOLD_JSON, encoding="utf-8") as f:
        return json.load(f)


class MeicaiParse(unittest.TestCase):
    def test_gold_rows_qty_amt_date_paid(self):
        parsed = core.parse_meicai_text(_ocr_text())
        gold = _gold()
        self.assertTrue(core.is_meicai_ocr(_ocr_text()))
        self.assertEqual(parsed["date"], gold["delivery_date"])
        self.assertEqual(parsed["paid"], gold["paid"])
        self.assertEqual(parsed["order_total"], gold["paid"])
        self.assertEqual(parsed["order_ids"], gold["order_ids"])
        self.assertIn("3002160333", parsed["note"])
        self.assertIn("3002160335", parsed["note"])
        rows = gold["日物料采购_复核后期望行"]
        self.assertEqual(len(parsed["items"]), len(rows))
        items, _stat = core.derive_price(parsed["items"])
        total = 0.0
        for got, want in zip(items, rows):
            self.assertEqual(got["品名"], want["name"])
            self.assertEqual(got["单位"], want["unit"])
            self.assertAlmostEqual(got["数量"], want["qty"], places=2)
            self.assertAlmostEqual(got["金额"], want["amt"], places=2)
            self.assertAlmostEqual(got["单价"], round(want["price"], 2), places=2)
            self.assertFalse(got.get("实收缺失"))
            self.assertFalse(got.get("员工区"))
            self.assertEqual(got.get("渠道"), "美菜")
            total += got["金额"]
        self.assertAlmostEqual(round(total, 2), gold["paid"], places=2)
        self.assertAlmostEqual(round(total, 2) - gold["paid"], gold["员工餐"]["amt"], places=2)
        names = [it["品名"] for it in items]
        self.assertNotIn("直降", names)
        self.assertFalse(any("欠款" in n for n in names))
        self.assertFalse(any("补差" in n or n.startswith("+") for n in names))
        # 实收，不是下单 3.00 斤 / 19.08 元
        self.assertAlmostEqual(items[0]["数量"], 3.08, places=2)
        self.assertAlmostEqual(items[0]["金额"], 19.59, places=2)
        self.assertNotAlmostEqual(items[0]["金额"], 19.08, places=2)

    def test_handwriting_html_is_not_meicai(self):
        with open(HAND_TXT, encoding="utf-8") as f:
            text = f.read()
        self.assertFalse(core.is_meicai_ocr(text))
        items = core.parse_line_items(core.parse_table(text))
        self.assertGreater(len(items), 0)
        self.assertEqual(core.parse_meicai_text(text)["items"], [])

    def test_missing_receive_line_does_not_copy_order_qty(self):
        text = "\n".join([
            "美菜",
            "送达时间 2026-09-10 06:00:00",
            "￥10.00",
            "下单实付￥9.00+需补交金额￥1.00",
            "订单号：3002160999 复制",
            "小葱 优质 水洗 斤 ￥10.00/2.00斤",
            "订 ￥10.00/2.00斤",
            "付款信息",
            "商品多退少补合计￥10.00",
            "历史欠款 ￥3.00",
            "直降 -￥1.00",
        ])
        parsed = core.parse_meicai_text(text)
        self.assertEqual(len(parsed["items"]), 1)
        it = parsed["items"][0]
        self.assertEqual(it["品名"], "小葱")
        self.assertIsNone(it["数量"])
        self.assertIsNone(it["金额"])
        self.assertTrue(it["实收缺失"])
        self.assertEqual(parsed["paid"], 10.0)
        self.assertTrue(any("实收" in s for s in parsed["suspects"]))
        self.assertNotEqual(parsed["paid"], 9.0)
        self.assertNotEqual(parsed["paid"], 12.0)

    def test_day_merge_concat_sum_and_conflict(self):
        a = core.parse_meicai_text(_ocr_text())
        b = {
            "date": "2026-09-09",
            "paid": 20.0,
            "order_ids": ["3002160336"],
            "items": [{
                "品名": "小葱", "数量": 1.0, "金额": 7.0, "单位": "斤",
                "渠道": "美菜", "订单号": "3002160336", "实收缺失": False,
            }],
            "suspects": [],
        }
        merged = core.merge_meicai_docs([a, b])
        self.assertEqual(len(merged), 1)
        doc = merged[0]
        # 9/9 里已有小葱 5/35.75，另一张再来 1/7，按日合并后再按品名+单位加总
        scallions = [it for it in doc["items"] if it["品名"] == "小葱"]
        self.assertEqual(len(scallions), 1)
        self.assertEqual(scallions[0]["单位"], "斤")
        self.assertAlmostEqual(scallions[0]["数量"], 6.0, places=2)
        self.assertAlmostEqual(scallions[0]["金额"], 42.75, places=2)
        self.assertFalse(scallions[0].get("实收缺失"))
        self.assertEqual(len(doc["items"]), len(a["items"]))
        self.assertAlmostEqual(doc["paid"], 183.25, places=2)
        self.assertEqual(doc["order_ids"], a["order_ids"] + ["3002160336"])

        again = core.merge_meicai_docs([a, a])
        self.assertEqual(len(again[0]["items"]), len(a["items"]))
        self.assertAlmostEqual(again[0]["paid"], a["paid"], places=2)

        conflict = {
            "date": "2026-09-09",
            "paid": 5.0,
            "order_ids": ["3002160333"],
            "items": [{
                "品名": "多肉排骨粒", "数量": 1.0, "金额": 5.0, "单位": "斤",
                "渠道": "美菜", "订单号": "3002160333",
            }],
            "suspects": [],
        }
        bad = core.merge_meicai_docs([a, conflict])
        self.assertEqual(len(bad[0]["items"]), len(a["items"]) + 1)
        self.assertTrue(any("3002160333" in s and "不一致" in s for s in bad[0]["suspects"]))
        self.assertEqual(bad[0]["order_ids"].count("3002160333"), 1)

        other_day = {
            "date": "2026-09-10",
            "paid": 1.0,
            "order_ids": ["9"],
            "items": [{"品名": "盐", "数量": 1, "金额": 1, "订单号": "9"}],
            "suspects": [],
        }
        split = core.merge_meicai_docs([a, other_day])
        self.assertEqual([d["date"] for d in split], ["2026-09-09", "2026-09-10"])

    def test_same_name_unit_merges_received_lines(self):
        """两行小葱（香葱）普通 3/26.91 + 5/44.40 → 一行 8.00 斤 / 71.31。"""
        text = "\n".join([
            "美菜",
            "送达时间 2026-09-03 06:30:00",
            "￥71.31",
            "订单号：2997320402 复制",
            "小葱（香葱）普通 斤 ￥26.91/3.00斤",
            "订 ￥26.91/3.00斤",
            "收 ￥26.91/3.00斤",
            "香菜 优质 斤 ￥19.64/3.04斤",
            "订 ￥19.64/3.04斤",
            "收 ￥19.64/3.04斤",
            "小葱（香葱）普通 斤 ￥44.40/5.00斤",
            "订 ￥44.40/5.00斤",
            "收 ￥44.40/5.00斤",
            "付款信息",
            "商品多退少补合计￥90.95",
        ])
        parsed = core.parse_meicai_text(text)
        self.assertEqual(
            [it["品名"] for it in parsed["items"]],
            ["小葱（香葱）普通", "香菜"],
        )
        scallion = parsed["items"][0]
        self.assertEqual(scallion["单位"], "斤")
        self.assertAlmostEqual(scallion["数量"], 8.0, places=2)
        self.assertAlmostEqual(scallion["金额"], 71.31, places=2)
        self.assertAlmostEqual(scallion["单价"], 8.91, places=2)
        self.assertFalse(scallion["实收缺失"])
        self.assertAlmostEqual(parsed["items"][1]["数量"], 3.04, places=2)
        self.assertAlmostEqual(parsed["items"][1]["金额"], 19.64, places=2)

    def test_same_name_different_unit_stays_split(self):
        text = "\n".join([
            "美菜",
            "送达时间 2026-09-03 06:30:00",
            "订单号：2997320402 复制",
            "小葱 优质 水洗 斤 ￥26.91/3.00斤",
            "收 ￥26.91/3.00斤",
            "小葱 优质 水洗 托 ￥10.00/1.00托",
            "收 ￥10.00/1.00托",
            "付款信息",
            "商品多退少补合计￥36.91",
        ])
        parsed = core.parse_meicai_text(text)
        self.assertEqual(
            [(it["品名"], it["单位"], it["数量"], it["金额"]) for it in parsed["items"]],
            [("小葱", "斤", 3.0, 26.91), ("小葱", "托", 1.0, 10.0)],
        )

    def test_empty_qty_amt_does_not_invent_numbers(self):
        """两边都没有实收：合成一行，数量和金额仍空。只有一边有数字就只加这一边。"""
        both_missing = "\n".join([
            "美菜",
            "送达时间 2026-09-03 06:30:00",
            "订单号：2997320402 复制",
            "小葱 优质 水洗 斤 ￥26.00/3.00斤",
            "订 ￥26.00/3.00斤",
            "小葱 优质 水洗 斤 ￥44.00/5.00斤",
            "订 ￥44.00/5.00斤",
            "付款信息",
            "商品多退少补合计￥71.31",
        ])
        parsed = core.parse_meicai_text(both_missing)
        scallions = [it for it in parsed["items"] if it["品名"] == "小葱"]
        self.assertEqual(len(scallions), 1)
        self.assertIsNone(scallions[0]["数量"])
        self.assertIsNone(scallions[0]["金额"])
        self.assertTrue(scallions[0]["实收缺失"])
        self.assertNotEqual(scallions[0]["数量"], 0)
        self.assertNotEqual(scallions[0]["数量"], 8)
        self.assertTrue(any("实收" in s for s in parsed["suspects"]))

        one_side = "\n".join([
            "美菜",
            "送达时间 2026-09-03 06:30:00",
            "订单号：2997320402 复制",
            "小葱 优质 水洗 斤 ￥26.91/3.00斤",
            "订 ￥26.91/3.00斤",
            "收 ￥26.91/3.00斤",
            "小葱 优质 水洗 斤 ￥44.40/5.00斤",
            "订 ￥44.40/5.00斤",
            "付款信息",
            "商品多退少补合计￥71.31",
        ])
        mixed = core.parse_meicai_text(one_side)
        row = next(it for it in mixed["items"] if it["品名"] == "小葱")
        self.assertEqual(len([it for it in mixed["items"] if it["品名"] == "小葱"]), 1)
        self.assertAlmostEqual(row["数量"], 3.0, places=2)
        self.assertAlmostEqual(row["金额"], 26.91, places=2)
        self.assertFalse(row["实收缺失"])
        self.assertNotAlmostEqual(row["数量"], 8.0, places=2)
        self.assertNotAlmostEqual(row["金额"], 71.31, places=2)

    def test_day_merge_then_same_name_unit_again(self):
        """两张图各有一行小葱，按日拼完再按品名+单位加总。复核单只留一行。"""
        def part(oid, qty, amt, filename):
            items = [{
                "品名": "小葱", "数量": qty, "金额": amt, "单位": "斤",
                "渠道": "美菜", "订单号": oid, "实收缺失": False,
            }]
            return {
                "filename": filename,
                "status": "pending",
                "ocr_done": True,
                "kind": "meicai",
                "date": "2026-09-03",
                "items": copy.deepcopy(items),
                "order_total": amt,
                "paid_total": amt,
                "order_ids": [oid],
                "note": "",
                "suspects": [],
                "meicai_part": {
                    "date": "2026-09-03",
                    "paid": amt,
                    "order_ids": [oid],
                    "items": items,
                    "note": "",
                    "suspects": [],
                },
            }

        invoices = [
            part("2997320400", 3.0, 26.91, "b.png"),
            part("2997320402", 5.0, 44.40, "a.png"),
        ]
        webapp.apply_meicai_day_merge(invoices)
        primary = next(inv for inv in invoices if inv["status"] != "merged")
        self.assertEqual(len(primary["items"]), 1)
        row = primary["items"][0]
        self.assertEqual(row["品名"], "小葱")
        self.assertEqual(row["单位"], "斤")
        self.assertAlmostEqual(row["数量"], 8.0, places=2)
        self.assertAlmostEqual(row["金额"], 71.31, places=2)
        self.assertAlmostEqual(row["单价"], 8.91, places=2)
        self.assertFalse(row.get("实收缺失"))
        self.assertEqual(sum(1 for inv in invoices if inv["status"] == "merged"), 1)

    def test_identical_order_lines_are_not_summed_twice(self):
        """同一订单号明细一致时不把行再加一遍。重叠截图不能把 3.08 斤加成 6.16。"""
        a = core.parse_meicai_text(_ocr_text())
        repeated_items = [it for it in a["items"] if it.get("订单号") == "3002160333"]
        repeated = {
            "date": a["date"],
            "paid": 10.0,
            "order_ids": ["3002160333", "3002160999"],
            "items": repeated_items + [{
                "品名": "生菜", "数量": 2.0, "金额": 10.0, "单位": "斤",
                "渠道": "美菜", "订单号": "3002160999", "实收缺失": False,
            }],
            "suspects": [],
        }
        merged = core.merge_meicai_docs([a, repeated])[0]
        rib = next(it for it in merged["items"] if it["品名"] == "多肉排骨粒")
        self.assertAlmostEqual(rib["数量"], 3.08, places=2)
        self.assertAlmostEqual(rib["金额"], 19.59, places=2)
        lettuce = [it for it in merged["items"] if it["品名"] == "生菜"]
        self.assertEqual(len(lettuce), 1)
        self.assertAlmostEqual(lettuce[0]["数量"], 2.0, places=2)
        self.assertFalse(any("不一致" in s for s in merged["suspects"]))

    def test_handwriting_same_name_is_not_merged(self):
        html = "\n".join([
            "<table><tr><td>编号</td><td>商品名称</td><td>单位</td><td>数量</td><td>单价</td><td>金额</td><td>备注</td></tr>",
            "<tr><td>1</td><td>小葱</td><td>斤</td><td>3</td><td>8.97</td><td>26.91</td><td></td></tr>",
            "<tr><td>2</td><td>小葱</td><td>斤</td><td>5</td><td>8.88</td><td>44.40</td><td></td></tr>",
            "</table>",
        ])
        self.assertFalse(core.is_meicai_ocr(html))
        items = core.parse_line_items(core.parse_table(html))
        self.assertEqual([it["品名"] for it in items], ["小葱", "小葱"])
        self.assertAlmostEqual(items[0]["数量"], 3.0, places=2)
        self.assertAlmostEqual(items[1]["数量"], 5.0, places=2)
        self.assertAlmostEqual(items[0]["金额"], 26.91, places=2)
        self.assertAlmostEqual(items[1]["金额"], 44.40, places=2)


class MeicaiRoute(unittest.TestCase):
    def setUp(self):
        self._call = webapp.core.call_ocr
        self._tiled = webapp.core.call_ocr_tiled
        self._state = copy.deepcopy(webapp._state)
        self._paths = dict(webapp._paths)

    def tearDown(self):
        webapp.core.call_ocr = self._call
        webapp.core.call_ocr_tiled = self._tiled
        webapp._state.clear()
        webapp._state.update(self._state)
        webapp._paths.clear()
        webapp._paths.update(self._paths)

    def test_ocr_one_routes_meicai_and_skips_enhance(self):
        calls = []

        def fake(path, timeout=600, enhance=False):
            calls.append(enhance)
            return _ocr_text()

        webapp.core.call_ocr = fake
        res = webapp.ocr_one("meicai_2026-09-09.png")
        self.assertEqual(calls, [False])
        self.assertEqual(res["kind"], "meicai")
        self.assertEqual(res["date"], "2026-09-09")
        self.assertEqual(res["order_total"], 163.25)
        self.assertEqual(res["paid_total"], 163.25)
        self.assertEqual(len(res["items"]), 7)
        gold = _gold()["日物料采购_复核后期望行"]
        for got, want in zip(res["items"], gold):
            self.assertEqual(got.get("原品名") or got["品名"], want["name"])
            self.assertAlmostEqual(got["数量"], want["qty"], places=2)
            self.assertAlmostEqual(got["金额"], want["amt"], places=2)
            self.assertNotEqual(got.get("归属"), "员工餐")
        self.assertEqual(res["meicai_part"]["order_ids"], _gold()["order_ids"])

    def test_day_merge_on_invoices_and_land_trace(self):
        text = _ocr_text()
        first = core.parse_meicai_text(text)
        second = {
            "date": "2026-09-09",
            "paid": 20.0,
            "order_ids": ["3002160999"],
            "note": "订单号 3002160999",
            "items": [{
                "品名": "生菜", "数量": 2.0, "金额": 20.0, "单位": "斤",
                "渠道": "美菜", "订单号": "3002160999", "实收缺失": False, "日期": "2026-09-09",
            }],
            "suspects": [],
        }
        parsed = [first, second]
        invoices = []
        for i, doc in enumerate(parsed):
            items, _ = core.derive_price(doc["items"])
            invoices.append({
                "filename": f"shot-{i}.png",
                "status": "pending",
                "ocr_done": True,
                "kind": "meicai",
                "date": doc["date"],
                "items": items,
                "order_total": doc["paid"],
                "paid_total": doc["paid"],
                "order_ids": doc["order_ids"],
                "note": doc["note"],
                "suspects": [],
                "meicai_part": {
                    "date": doc["date"],
                    "paid": doc["paid"],
                    "order_ids": doc["order_ids"],
                    "items": copy.deepcopy(items),
                    "note": doc["note"],
                    "suspects": doc["suspects"],
                },
            })
        webapp.apply_meicai_day_merge(invoices)
        self.assertEqual(invoices[0]["status"], "pending")
        self.assertEqual(len(invoices[0]["items"]), 8)
        self.assertEqual(invoices[0]["items"][-1]["品名"], "生菜")
        self.assertEqual(invoices[1]["status"], "merged")
        self.assertEqual(invoices[1]["items"], [])
        self.assertIn("3002160333", invoices[0]["note"])
        self.assertIn("3002160999", invoices[0]["order_ids"])
        self.assertEqual(invoices[0]["order_ids"].count("3002160333"), 1)
        self.assertAlmostEqual(invoices[0]["paid_total"], 183.25, places=2)

        tmp = tempfile.TemporaryDirectory()
        self.addCleanup(tmp.cleanup)
        webapp._paths["output_dir"] = tmp.name
        webapp.append_land(
            "shot-0.png", "2026-09-09", "蔬菜",
            invoices[0]["total"], invoices[0]["paid_total"], invoices[0]["items"],
            note=invoices[0]["note"],
        )
        wb = load_workbook(os.path.join(tmp.name, "采购入库_落地.xlsx"))
        ws = wb.active
        self.assertEqual(ws.cell(1, 9).value, "实付合计")
        self.assertEqual(ws.cell(1, 13).value, "溯源")
        self.assertEqual(ws.cell(2, 3).value, "多肉排骨粒")
        self.assertAlmostEqual(ws.cell(2, 4).value, 3.08, places=2)
        self.assertAlmostEqual(ws.cell(2, 6).value, 19.59, places=2)
        self.assertIn("3002160333", str(ws.cell(2, 13).value))
        self.assertEqual(ws.max_row, 9)  # header + 7 goods + 并入的生菜
        wb.close()

    def test_confirm_rejects_merged_screenshot(self):
        webapp._state["invoices"] = [{
            "filename": "shot-1.png",
            "status": "merged",
            "ocr_done": True,
            "kind": "meicai",
            "date": "2026-09-09",
            "items": [],
            "note": "已并入 shot-0.png",
        }]
        with self.assertRaises(HTTPException) as ctx:
            webapp.confirm("shot-1.png", webapp.ConfirmPayload(
                date="2026-09-09", category="蔬菜",
                items=[{"品名": "小葱", "数量": 1, "金额": 1}],
                paid_total=1,
            ))
        self.assertEqual(ctx.exception.status_code, 400)
        self.assertIn("并入", ctx.exception.detail)


class LongScreenshot(unittest.TestCase):
    def test_slices_cover_ends_without_exceeding_edge(self):
        img = Image.new("RGB", (1080, 4926), "white")
        img.putpixel((0, 0), (255, 0, 0))
        img.putpixel((10, 4925), (0, 255, 0))
        slices = list(core.iter_vertical_slices(img, max_edge=2048, overlap=0.15))
        self.assertGreaterEqual(len(slices), 3)
        self.assertTrue(all(max(s.size) <= 2048 for s in slices))
        self.assertTrue(any(s.getpixel((0, 0)) == (255, 0, 0) for s in slices))
        self.assertTrue(any(s.getpixel((10, s.size[1] - 1)) == (0, 255, 0) for s in slices))

    def test_join_drops_exact_overlap(self):
        joined = core.join_ocr_slices([
            "订单号：1\n收 ￥1.00/1.00斤\n鸡毛菜",
            "鸡毛菜\n收 ￥2.00/2.00斤",
        ])
        self.assertEqual(joined, "订单号：1\n收 ￥1.00/1.00斤\n鸡毛菜\n收 ￥2.00/2.00斤")

    def test_geometry_gate(self):
        tmp = tempfile.TemporaryDirectory()
        self.addCleanup(tmp.cleanup)
        tall = os.path.join(tmp.name, "tall.png")
        short = os.path.join(tmp.name, "short.png")
        Image.new("RGB", (1080, 4926), "white").save(tall)
        Image.new("RGB", (1600, 2000), "white").save(short)
        self.assertTrue(core.looks_like_long_screenshot(tall))
        self.assertFalse(core.looks_like_long_screenshot(short))
        self.assertFalse(core.looks_like_long_screenshot(os.path.join(tmp.name, "missing.png")))

    def test_in_limit_frame_stays_original_photo_still_capped(self):
        """上限内的图不缩小。手写单那种大照片仍把最长边收到 2048。"""
        tmp = tempfile.TemporaryDirectory()
        self.addCleanup(tmp.cleanup)
        small = os.path.join(tmp.name, "small.jpg")
        photo = os.path.join(tmp.name, "photo.jpg")
        Image.new("RGB", (1080, 1800), "white").save(small, quality=95)
        Image.new("RGB", (3024, 4032), "white").save(photo, quality=85)
        small_buf, _ = core.prepare_image(small)
        photo_buf, _ = core.prepare_image(photo)
        self.assertEqual(Image.open(small_buf).size, (1080, 1800))
        photo_size = Image.open(photo_buf).size
        self.assertEqual(max(photo_size), 2048)
        self.assertLess(photo_size[0], 3024)
        self.assertLess(photo_size[1], 4032)

    def test_phone_long_slices_keep_native_width_and_cover_rows(self):
        w, h = 1080, 6000
        img = Image.new("RGB", (w, h), "white")
        img.putpixel((0, 0), (255, 0, 0))
        img.putpixel((20, h - 1), (0, 255, 0))
        slices = list(core.iter_ocr_slices(img))
        self.assertGreaterEqual(len(slices), 5)
        self.assertLess(core.OCR_SLICE_HEIGHT, core.MAX_SIZE)
        self.assertTrue(all(s.size[0] == w for s in slices))
        self.assertTrue(all(s.size[1] <= core.OCR_SLICE_HEIGHT for s in slices))
        self.assertTrue(all(max(s.size) <= core.MAX_SIZE for s in slices))
        self.assertTrue(any(s.getpixel((0, 0)) == (255, 0, 0) for s in slices))
        self.assertTrue(any(s.getpixel((20, s.size[1] - 1)) == (0, 255, 0) for s in slices))
        # 整图路径仍会把 1080×6000 压到大约 370 宽。分条不能跟着压。
        ratio = core.full_frame_shrink_ratio(w, h)
        self.assertLessEqual(ratio, core.CRUSHED_FULL_FRAME_RATIO)
        self.assertLess(int(w * ratio), 400)
        encoded, _ = core._image_to_jpeg_buf(slices[0], "slice-0.jpg")
        self.assertEqual(Image.open(encoded).size, slices[0].size)

        boxes = core.ocr_slice_boxes(w, h)
        self.assertEqual(len(boxes), len(slices))
        self.assertEqual(boxes[0][0], 0)
        self.assertEqual(boxes[-1][1], h)
        self.assertTrue(all(b - a <= core.OCR_SLICE_HEIGHT for a, b in boxes))
        row = max(1, int(round(w * core.SLICE_ROW_OVERLAP_RATIO)))
        for (a0, a1), (b0, b1) in zip(boxes, boxes[1:]):
            self.assertGreaterEqual(a1 - b0, row)
        for y0 in range(0, h - row + 1, 37):
            y1 = y0 + row
            self.assertTrue(
                any(a <= y0 and y1 <= b for a, b in boxes),
                "row %s:%s cut across %s" % (y0, y1, boxes),
            )

    def test_wide_long_image_caps_width_not_the_long_edge(self):
        # 3000×5000 若按最长边收到 2048，宽只剩约 1228。分条只把宽收到 2048。
        img = Image.new("RGB", (3000, 5000), "white")
        slices = list(core.iter_ocr_slices(img))
        self.assertGreaterEqual(len(slices), 2)
        self.assertTrue(all(s.size[0] == core.MAX_SIZE for s in slices))
        self.assertTrue(all(s.size[1] <= core.OCR_SLICE_HEIGHT for s in slices))
        self.assertTrue(all(max(s.size) <= core.MAX_SIZE for s in slices))
        crushed_w = int(3000 * (core.MAX_SIZE / 5000.0))
        self.assertGreater(slices[0].size[0], crushed_w)
        # 编码不再按最长边把这条压回去。宽已是上限时尺寸不变。
        encoded, _ = core._image_to_jpeg_buf(slices[0], "slice-0.jpg")
        self.assertEqual(Image.open(encoded).size, slices[0].size)

    def test_slice_jpeg_shrinks_width_only(self):
        """分条编码只缩宽。高过上限的窄条不能按最长边把宽一起压瘦。"""
        tall = Image.new("RGB", (1080, 2500), "white")
        tall_buf, _ = core._image_to_jpeg_buf(tall, "tall.jpg")
        self.assertEqual(Image.open(tall_buf).size, (1080, 2500))
        wide = Image.new("RGB", (3000, 4000), "white")
        wide_buf, _ = core._image_to_jpeg_buf(wide, "wide.jpg")
        opened = Image.open(wide_buf)
        fitted_h = int(4000 * (core.MAX_SIZE / 3000.0))
        self.assertEqual(opened.size, (core.MAX_SIZE, fitted_h))
        self.assertGreater(opened.size[1], core.MAX_SIZE)
        longest = int(3000 * (core.MAX_SIZE / 4000.0))
        self.assertGreater(opened.size[0], longest)

    def test_ocr_one_uses_tiled_text_for_long_image(self):
        tmp = tempfile.TemporaryDirectory()
        self.addCleanup(tmp.cleanup)
        path_dir = tmp.name
        Image.new("RGB", (1080, 3000), "white").save(os.path.join(path_dir, "long.png"))
        old_paths = dict(webapp._paths)
        old_call = webapp.core.call_ocr
        old_tiled = webapp.core.call_ocr_tiled
        webapp._paths["input_dir"] = path_dir
        seen = {"tiled": 0, "full": 0}

        def tiled(path, timeout=600, overlap=0.15):
            seen["tiled"] += 1
            return _ocr_text()

        def full(path, timeout=600, enhance=False):
            seen["full"] += 1
            return "<table><tr><td>1</td><td>白菜</td><td>1</td><td>2</td><td>2</td></tr></table>"

        webapp.core.call_ocr_tiled = tiled
        webapp.core.call_ocr = full
        try:
            res = webapp.ocr_one("long.png")
        finally:
            webapp.core.call_ocr = old_call
            webapp.core.call_ocr_tiled = old_tiled
            webapp._paths.clear()
            webapp._paths.update(old_paths)
        self.assertEqual(seen["tiled"], 1)
        # 长图会再打一次整图，但整图不是美菜实收，分数更低，仍用分条原文。
        self.assertEqual(seen["full"], 1)
        self.assertEqual(res["kind"], "meicai")
        self.assertEqual(len(res["items"]), 7)
        self.assertAlmostEqual(res["items"][0]["数量"], 3.08, places=2)

    def test_ocr_one_prefers_full_when_it_has_more_recv(self):
        """分条没有可用「收」、整图「收」多出足够行数时，用整图原文。不把订货数写成实收。"""
        tmp = tempfile.TemporaryDirectory()
        self.addCleanup(tmp.cleanup)
        path_dir = tmp.name
        Image.new("RGB", (1080, 3000), "white").save(os.path.join(path_dir, "long.png"))
        old_paths = dict(webapp._paths)
        old_call = webapp.core.call_ocr
        old_tiled = webapp.core.call_ocr_tiled
        webapp._paths["input_dir"] = path_dir

        def tiled(path, timeout=600, overlap=0.15):
            return "\n".join([
                "美菜",
                "送达时间 2026-09-09 06:35:03",
                "订单号：3002160333 复制",
                "(56,141),(117,169)已完成",
                "美菜配送-已送达",
                "鸡毛菜 斤",
                "￥14.55/3.00斤",
                "订",
                "￥14.55/3.00斤",
                "付款信息",
                "商品多退少补合计￥163.25",
            ])

        def full(path, timeout=600, enhance=False):
            return _ocr_text()

        webapp.core.call_ocr_tiled = tiled
        webapp.core.call_ocr = full
        try:
            res = webapp.ocr_one("long.png")
        finally:
            webapp.core.call_ocr = old_call
            webapp.core.call_ocr_tiled = old_tiled
            webapp._paths.clear()
            webapp._paths.update(old_paths)
        self.assertEqual(res["kind"], "meicai")
        self.assertEqual(len(res["items"]), 7)
        self.assertAlmostEqual(res["items"][0]["数量"], 3.08, places=2)
        self.assertAlmostEqual(res["items"][0]["金额"], 19.59, places=2)

    def _ocr_long(self, size, tiled_text, full_text):
        tmp = tempfile.TemporaryDirectory()
        self.addCleanup(tmp.cleanup)
        Image.new("RGB", size, "white").save(os.path.join(tmp.name, "long.png"))
        old_paths = dict(webapp._paths)
        old_call = webapp.core.call_ocr
        old_tiled = webapp.core.call_ocr_tiled
        webapp._paths["input_dir"] = tmp.name
        seen = {"tiled": 0, "full": 0}

        def tiled(path, timeout=600, overlap=0.15):
            seen["tiled"] += 1
            return tiled_text

        def full(path, timeout=600, enhance=False):
            seen["full"] += 1
            return full_text

        webapp.core.call_ocr_tiled = tiled
        webapp.core.call_ocr = full
        try:
            res = webapp.ocr_one("long.png")
        finally:
            webapp.core.call_ocr = old_call
            webapp.core.call_ocr_tiled = old_tiled
            webapp._paths.clear()
            webapp._paths.update(old_paths)
        return res, seen

    def test_crushed_full_does_not_beat_tiled_with_as_many_recv(self):
        """1080×6000 的整图会被压瘦。分条「收」不少于整图时，不用整图多出来的品名行。"""
        tiled_text = "\n".join([
            "美菜",
            "送达时间 2026-09-09 06:35:03",
            "订单号：3002160333 复制",
            "小葱 斤 ￥10.00/1.00斤",
            "收 ￥10.00/1.00斤",
            "付款信息",
            "商品多退少补合计￥10.00",
        ])
        full_text = "\n".join([
            "美菜",
            "送达时间 2026-09-09 06:35:03",
            "订单号：3002160333 复制",
            "小葱 斤 ￥10.00/1.00斤",
            "收 ￥10.00/1.00斤",
            "鸡毛菜 斤 ￥6.00/1.00斤",
            "付款信息",
            "商品多退少补合计￥10.00",
        ])
        self.assertGreater(core.meicai_ocr_quality(full_text), core.meicai_ocr_quality(tiled_text))
        self.assertEqual(core.meicai_recv_count(tiled_text), core.meicai_recv_count(full_text))
        res, seen = self._ocr_long((1080, 6000), tiled_text, full_text)
        self.assertEqual(seen["tiled"], 1)
        self.assertEqual(seen["full"], 1)
        self.assertEqual([it["品名"] for it in res["items"]], ["小葱"])

    def test_one_extra_recv_does_not_let_shrunk_full_win(self):
        """压瘦的整图只多一行「收」时仍用分条。不拿糊掉的字盖过原宽切片。"""
        tiled_text = "\n".join([
            "美菜",
            "送达时间 2026-09-09 06:35:03",
            "订单号：3002160333 复制",
            "小葱 斤 ￥10.00/1.00斤",
            "收 ￥10.00/1.00斤",
            "付款信息",
            "商品多退少补合计￥16.00",
        ])
        full_text = "\n".join([
            "美菜",
            "送达时间 2026-09-09 06:35:03",
            "订单号：3002160333 复制",
            "小葱 斤 ￥10.00/1.00斤",
            "收 ￥10.00/1.00斤",
            "鸡毛菜 斤 ￥6.00/1.00斤",
            "收 ￥6.00/1.00斤",
            "付款信息",
            "商品多退少补合计￥16.00",
        ])
        self.assertFalse(core.shrunk_full_beats_tiled(tiled_text, full_text))
        res, seen = self._ocr_long((1080, 6000), tiled_text, full_text)
        self.assertEqual(seen["full"], 1)
        self.assertEqual([it["品名"] for it in res["items"]], ["小葱"])

    def test_full_wins_only_when_recv_margin_is_met(self):
        """分条「收」明显更少（少够 MEICAI_FULL_RECV_MARGIN 行）时才退回整图。"""
        tiled_rows = [
            "美菜",
            "送达时间 2026-09-09 06:35:03",
            "订单号：3002160333 复制",
            "小葱 斤 ￥10.00/1.00斤",
            "收 ￥10.00/1.00斤",
            "付款信息",
            "商品多退少补合计￥40.00",
        ]
        names = ["小葱", "鸡毛菜", "生菜", "香菜"]
        amounts = [(10.0, 1.0), (6.0, 1.0), (8.0, 2.0), (4.0, 1.0)]
        self.assertEqual(len(names), 1 + core.MEICAI_FULL_RECV_MARGIN)
        full_rows = [
            "美菜",
            "送达时间 2026-09-09 06:35:03",
            "订单号：3002160333 复制",
        ]
        for name, (amt, qty) in zip(names, amounts):
            full_rows.append("%s 斤 ￥%.2f/%.2f斤" % (name, amt, qty))
            full_rows.append("收 ￥%.2f/%.2f斤" % (amt, qty))
        full_rows.extend(["付款信息", "商品多退少补合计￥40.00"])
        tiled_text = "\n".join(tiled_rows)
        full_text = "\n".join(full_rows)
        self.assertTrue(core.shrunk_full_beats_tiled(tiled_text, full_text))
        # 只留前三件货：比分条多两行「收」，还没到门槛。
        short_full = "\n".join(full_rows[:3] + full_rows[3:9] + full_rows[-2:])
        self.assertEqual(core.meicai_recv_count(short_full), core.meicai_recv_count(tiled_text) + 2)
        self.assertFalse(core.shrunk_full_beats_tiled(tiled_text, short_full))
        res, seen = self._ocr_long((1080, 6000), tiled_text, full_text)
        self.assertEqual(seen["tiled"], 1)
        self.assertEqual(seen["full"], 1)
        self.assertEqual([it["品名"] for it in res["items"]], names)
        self.assertAlmostEqual(res["items"][1]["数量"], 1.0, places=2)
        self.assertAlmostEqual(res["items"][1]["金额"], 6.0, places=2)

    def test_mild_shrink_does_not_win_on_extra_goods_lines(self):
        """1080×3000 也会把长边收到 2048。整图多一行品名、没有多出来的「收」，仍用分条。"""
        tiled_text = "\n".join([
            "美菜",
            "送达时间 2026-09-09 06:35:03",
            "订单号：3002160333 复制",
            "小葱 斤 ￥10.00/1.00斤",
            "收 ￥10.00/1.00斤",
            "付款信息",
            "商品多退少补合计￥10.00",
        ])
        full_text = "\n".join([
            "美菜",
            "送达时间 2026-09-09 06:35:03",
            "订单号：3002160333 复制",
            "小葱 斤 ￥10.00/1.00斤",
            "收 ￥10.00/1.00斤",
            "鸡毛菜 斤 ￥6.00/1.00斤",
            "付款信息",
            "商品多退少补合计￥10.00",
        ])
        self.assertGreater(core.full_frame_shrink_ratio(1080, 3000), core.CRUSHED_FULL_FRAME_RATIO)
        self.assertGreater(core.meicai_ocr_quality(full_text), core.meicai_ocr_quality(tiled_text))
        self.assertFalse(core.shrunk_full_beats_tiled(tiled_text, full_text))
        res, _seen = self._ocr_long((1080, 3000), tiled_text, full_text)
        self.assertEqual([it["品名"] for it in res["items"]], ["小葱"])

    def test_long_handwriting_stays_on_tiled_text(self):
        """两份都不像美菜时不换文本。手写单仍走原来的表解析。"""
        tmp = tempfile.TemporaryDirectory()
        self.addCleanup(tmp.cleanup)
        path_dir = tmp.name
        Image.new("RGB", (1080, 3000), "white").save(os.path.join(path_dir, "long.png"))
        old_paths = dict(webapp._paths)
        old_call = webapp.core.call_ocr
        old_tiled = webapp.core.call_ocr_tiled
        webapp._paths["input_dir"] = path_dir
        with open(HAND_TXT, encoding="utf-8") as f:
            hand = f.read()

        def tiled(path, timeout=600, overlap=0.15):
            return hand

        def full(path, timeout=600, enhance=False):
            return "<table><tr><td>1</td><td>不该用这份</td><td>9</td><td>9</td><td>9</td></tr></table>"

        webapp.core.call_ocr_tiled = tiled
        webapp.core.call_ocr = full
        try:
            res = webapp.ocr_one("long.png")
        finally:
            webapp.core.call_ocr = old_call
            webapp.core.call_ocr_tiled = old_tiled
            webapp._paths.clear()
            webapp._paths.update(old_paths)
        self.assertEqual(res["kind"], "handwriting")
        names = [it["品名"] for it in res["items"]]
        self.assertIn("薄荷叶", names)
        self.assertNotIn("不该用这份", names)

    def test_ocr_one_merges_same_meicai_name_before_review(self):
        text = "\n".join([
            "美菜",
            "送达时间 2026-09-03 06:30:00",
            "订单号：2997320402 复制",
            "小葱（香葱）普通 斤 ￥26.91/3.00斤",
            "订 ￥26.91/3.00斤",
            "收 ￥26.91/3.00斤",
            "小葱（香葱）普通 斤 ￥44.40/5.00斤",
            "订 ￥44.40/5.00斤",
            "收 ￥44.40/5.00斤",
            "付款信息",
            "商品多退少补合计￥71.31",
        ])

        def fake(path, timeout=600, enhance=False):
            return text

        webapp.core.call_ocr = fake
        res = webapp.ocr_one("meicai_scallion.png")
        self.assertEqual(res["kind"], "meicai")
        self.assertEqual(len(res["items"]), 1)
        row = res["items"][0]
        self.assertEqual(row["单位"], "斤")
        self.assertAlmostEqual(row["数量"], 8.0, places=2)
        self.assertAlmostEqual(row["金额"], 71.31, places=2)
        self.assertAlmostEqual(row["单价"], 8.91, places=2)
        self.assertEqual(len(res["meicai_part"]["items"]), 1)

    def test_ocr_one_handwriting_keeps_duplicate_names(self):
        html = "\n".join([
            "<table><tr><td>编号</td><td>商品名称</td><td>单位</td><td>数量</td><td>单价</td><td>金额</td><td>备注</td></tr>",
            "<tr><td>1</td><td>小葱</td><td>斤</td><td>3</td><td>8.97</td><td>26.91</td><td></td></tr>",
            "<tr><td>2</td><td>小葱</td><td>斤</td><td>5</td><td>8.88</td><td>44.40</td><td></td></tr>",
            "</table>",
        ])

        def fake(path, timeout=600, enhance=False):
            return html

        webapp.core.call_ocr = fake
        res = webapp.ocr_one("hand_scallion.png")
        self.assertEqual(res["kind"], "handwriting")
        self.assertEqual([it["品名"] for it in res["items"]], ["小葱", "小葱"])
        self.assertAlmostEqual(res["items"][0]["数量"], 3.0, places=2)
        self.assertAlmostEqual(res["items"][1]["数量"], 5.0, places=2)
        self.assertAlmostEqual(res["items"][0]["金额"], 26.91, places=2)
        self.assertAlmostEqual(res["items"][1]["金额"], 44.40, places=2)


def _read_json(path):
    with open(path, encoding="utf-8") as f:
        return json.load(f)


def _shot_gold(name):
    return _read_json(SHOT_GOLD[name])


def _doc_from_gold(gold):
    items = []
    for row in gold["日物料采购_复核后期望行"]:
        items.append({
            "品名": row["name"],
            "数量": row["qty"],
            "金额": row["amt"],
            "单位": row["unit"],
            "渠道": "美菜",
            "订单号": "",
            "实收缺失": False,
        })
    return {
        "date": gold["delivery_date"],
        "paid": gold["paid"],
        "order_ids": list(gold["order_ids"]),
        "items": items,
        "suspects": [],
    }


def _assert_filled_matches_gold(test, parsed, gold):
    """「收」行齐全时，解析结果应落在复核金标上（单价按分四舍五入）。"""
    test.assertEqual(parsed["date"], gold["delivery_date"])
    test.assertEqual(parsed["paid"], gold["paid"])
    test.assertEqual(parsed["order_ids"], gold["order_ids"])
    rows = gold["日物料采购_复核后期望行"]
    test.assertEqual(len(parsed["items"]), len(rows))
    items, _stat = core.derive_price(parsed["items"])
    total = 0.0
    for got, want in zip(items, rows):
        test.assertEqual(got["品名"], want["name"])
        test.assertEqual(got["单位"], want["unit"])
        test.assertAlmostEqual(got["数量"], want["qty"], places=2)
        test.assertAlmostEqual(got["金额"], want["amt"], places=2)
        test.assertFalse(got.get("实收缺失"))
        total += got["金额"]
    test.assertAlmostEqual(round(total, 2), gold["paid"], places=2)
    test.assertEqual(gold["员工餐"]["amt"], 0)
    test.assertAlmostEqual(round(total, 2) - gold["paid"], 0, places=2)


def _assert_missing_receive_not_invented(test, parsed, gold, text):
    """没有「收」的行不许填金标实收，也不许填标题上的订货数。"""
    test.assertEqual(parsed["date"], gold["delivery_date"])
    if re.search(r"商品多退少补合计\s*[￥¥]\s*" + re.escape(f"{gold['paid']:.2f}"), text):
        test.assertEqual(parsed["paid"], gold["paid"])
    recv_n = len(re.findall(r"(?:^|\n)\s*收\s*[￥¥]", text))
    filled = [it for it in parsed["items"] if it.get("数量") is not None or it.get("金额") is not None]
    test.assertLessEqual(len(filled), recv_n)
    for it in parsed["items"]:
        if it.get("实收缺失") or recv_n == 0:
            test.assertIsNone(it.get("数量"))
            test.assertIsNone(it.get("金额"))
            test.assertTrue(it.get("实收缺失"))
    for oid in parsed["order_ids"]:
        test.assertIn(oid, text)
    if recv_n == 0:
        test.assertTrue(any("实收" in s for s in parsed["suspects"]))
        gold_amts = {round(float(row["amt"]), 2) for row in gold["日物料采购_复核后期望行"]}
        for it in parsed["items"]:
            test.assertNotIn(it.get("金额"), gold_amts)


class ShotGoldTargets(unittest.TestCase):
    """Shot1 / Shot2 是复核落地金标。9/9 OCR 文本仍是解析必须通过的那条。"""

    def test_fixtures_match_approved_landing_totals(self):
        shot1 = _shot_gold("shot1")
        shot2 = _shot_gold("shot2")
        for gold, day, paid, n_orders in (
            (shot1, "2026-09-02", 258.23, 4),
            (shot2, "2026-09-03", 258.99, 3),
        ):
            rows = gold["日物料采购_复核后期望行"]
            self.assertEqual(gold["delivery_date"], day)
            self.assertEqual(gold["paid"], paid)
            self.assertEqual(gold["员工餐"]["amt"], 0)
            self.assertTrue(gold["paid_matches_line_sum"])
            self.assertAlmostEqual(sum(r["amt"] for r in rows), paid, places=2)
            self.assertEqual(len(gold["order_ids"]), n_orders)
            self.assertEqual(len(gold["order_ids"]), len(set(gold["order_ids"])))
        # Shot2 两行小葱落地成一行：3/26.91 + 5/44.40 → 8/71.31，总行数少 1
        self.assertEqual(len(shot2["日物料采购_复核后期望行"]), 9)
        scallions = [r for r in shot2["日物料采购_复核后期望行"] if r["name"] == "小葱"]
        self.assertEqual(len(scallions), 1)
        self.assertEqual(scallions[0]["unit"], "斤")
        self.assertAlmostEqual(scallions[0]["qty"], 8.0, places=2)
        self.assertAlmostEqual(scallions[0]["amt"], 71.31, places=2)
        shot1_scallions = [r for r in shot1["日物料采购_复核后期望行"] if r["name"] == "小葱"]
        self.assertEqual(len(shot1_scallions), 1)
        self.assertAlmostEqual(shot1_scallions[0]["qty"], 5.0, places=2)

    def test_different_delivery_dates_stay_separate_docs(self):
        merged = core.merge_meicai_docs([
            _doc_from_gold(_shot_gold("shot1")),
            _doc_from_gold(_shot_gold("shot2")),
        ])
        self.assertEqual([d["date"] for d in merged], ["2026-09-02", "2026-09-03"])
        self.assertAlmostEqual(merged[0]["paid"], 258.23, places=2)
        self.assertAlmostEqual(merged[1]["paid"], 258.99, places=2)
        self.assertEqual(len(merged[1]["items"]), 9)
        scallion = next(it for it in merged[1]["items"] if it["品名"] == "小葱")
        self.assertAlmostEqual(scallion["数量"], 8.0, places=2)
        self.assertAlmostEqual(scallion["金额"], 71.31, places=2)

    def test_shot2_weak_ocr_marks_receive_missing(self):
        """按 Shot2 金标备注构造的弱 OCR：没有「收」分行，数字是订货量，还漏了一个订单号。"""
        text = "\n".join([
            "美菜",
            "送达时间 2026-09-03 06:30:00",
            "￥258.99",
            "下单实付￥250.00",
            "订单号：2997320400 复制",
            "鸡毛菜 斤 ￥36.00/6.00斤",
            "订 ￥36.00/6.00斤",
            "黄瓜 二级 斤 ￥11.00/5.00斤",
            "订 ￥11.00/5.00斤",
            "多肉排骨粒 斤 ￥19.08/3.00斤",
            "订 ￥19.08/3.00斤",
            "鲜猪后腿瘦肉 纯瘦 斤 ￥48.00/5.00斤",
            "订 ￥48.00/5.00斤",
            "绿豆芽 斤 ￥3.00/3.00斤",
            "订 ￥3.00/3.00斤",
            "订单号：2997320402 复制",
            "小葱 优质 水洗 斤 ￥26.00/3.00斤",
            "订 ￥26.00/3.00斤",
            "小葱 优质 水洗 斤 ￥44.00/5.00斤",
            "订 ￥44.00/5.00斤",
            "香菜 斤 ￥19.00/3.00斤",
            "订 ￥19.00/3.00斤",
            "鲜鸡蛋 白壳 中码 托(30 个) ￥40.00/2.00托",
            "订 ￥40.00/2.00托",
            "鱼腥草 斤 ￥6.00/1.00斤",
            "订 ￥6.00/1.00斤",
            "付款信息",
            "直降 -￥5.00",
            "历史欠款 ￥2.42",
            "商品多退少补合计￥258.99",
        ])
        gold = _shot_gold("shot2")
        parsed = core.parse_meicai_text(text)
        self.assertTrue(core.is_meicai_ocr(text))
        _assert_missing_receive_not_invented(self, parsed, gold, text)
        self.assertEqual(parsed["order_ids"], ["2997320400", "2997320402"])
        self.assertNotIn("2997320401", parsed["order_ids"])
        names = [it["品名"] for it in parsed["items"]]
        # 两行都没有「收」，合成一行后仍是实收缺失，不能编出 8/71.31
        self.assertEqual(names.count("小葱"), 1)
        weak_scallion = next(it for it in parsed["items"] if it["品名"] == "小葱")
        self.assertIsNone(weak_scallion["数量"])
        self.assertIsNone(weak_scallion["金额"])
        self.assertTrue(weak_scallion["实收缺失"])
        self.assertIn("黄瓜 二级", names)
        self.assertIn("鲜鸡蛋 白壳", names)
        self.assertNotIn("直降", names)
        self.assertFalse(any("欠款" in n for n in names))
        # 订货数 5.00 / 11.00 不能当成黄瓜的实收（金标是 5.08 斤、11.84 元）
        cucumber = next(it for it in parsed["items"] if it["品名"] == "黄瓜 二级")
        self.assertIsNone(cucumber["数量"])
        self.assertIsNone(cucumber["金额"])

    def test_optional_shot_ocr_transcripts(self):
        """有 OCR 文本才和金标行对账。Shot2 若「收」行不够，只检查不编造。"""
        ran = False
        for name in ("shot1", "shot2"):
            path = SHOT_OCR[name]
            if not os.path.isfile(path):
                continue
            ran = True
            with open(path, encoding="utf-8") as f:
                text = f.read()
            gold = _shot_gold(name)
            parsed = core.parse_meicai_text(text)
            self.assertTrue(core.is_meicai_ocr(text), name)
            recv_n = len(re.findall(r"(?:^|\n)\s*收\s*[￥¥]", text))
            if recv_n >= len(gold["日物料采购_复核后期望行"]):
                _assert_filled_matches_gold(self, parsed, gold)
            else:
                _assert_missing_receive_not_invented(self, parsed, gold, text)
        if not ran:
            self.skipTest("shot1/shot2 的 OCR 文本还没放进 samples/gold/ocr/")

    def test_shot_misreads_land_on_gold_names(self):
        """错认文本里的别名和错挂收，落地后就是金标品名，不留促重行和空数量小葱。"""
        for name, banned in (
            ("shot1", ("白壳中码", "汽水 碳酸饮料", "汽水碳酸饮料")),
            ("shot2", ("促鸡毛菜", "促鱼腥草", "美菜鲜坊", "白壳中码")),
        ):
            text = _read_text(SHOT_OCR[name])
            gold = _shot_gold(name)
            parsed = core.parse_meicai_text(text)
            _assert_filled_matches_gold(self, parsed, gold)
            names = [it["品名"] for it in parsed["items"]]
            for word in banned:
                self.assertNotIn(word, "\n".join(names))
        shot2 = core.parse_meicai_text(_read_text(SHOT_OCR["shot2"]))
        scallion = next(it for it in shot2["items"] if it["品名"] == "小葱")
        self.assertAlmostEqual(scallion["数量"], 8.0, places=2)
        self.assertAlmostEqual(scallion["金额"], 71.31, places=2)
        self.assertFalse(scallion["实收缺失"])
        self.assertEqual([it["品名"] for it in shot2["items"]].count("鸡毛菜"), 1)
        self.assertEqual([it["品名"] for it in shot2["items"]].count("香菜"), 1)
        herb = next(it for it in shot2["items"] if it["品名"] == "鱼腥草")
        self.assertAlmostEqual(herb["数量"], 1.0, places=2)
        self.assertAlmostEqual(herb["金额"], 6.18, places=2)

    def test_shot_ocr_fixtures_keep_coriander_scallion_and_herb(self):
        """实机三条只认仓库里的 OCR 原文，不认另外编的价目表。

        Shot1：香菜 2 斤 / 13.70 在明细里，整单 6 行。
        Shot2：小葱合并 8 斤 / 71.31；香菜单独一行；鱼腥草是 1 斤 / 6.18，不是 4.04/25.82。
        """
        shot1 = core.parse_meicai_text(_read_text(SHOT_OCR["shot1"]))
        _assert_filled_matches_gold(self, shot1, _shot_gold("shot1"))
        coriander = next(it for it in shot1["items"] if it["品名"] == "香菜")
        self.assertAlmostEqual(coriander["数量"], 2.0, places=2)
        self.assertAlmostEqual(coriander["金额"], 13.70, places=2)
        self.assertEqual(len(shot1["items"]), 6)
        self.assertAlmostEqual(sum(it["金额"] for it in shot1["items"]), 258.23, places=2)

        shot2 = core.parse_meicai_text(_read_text(SHOT_OCR["shot2"]))
        _assert_filled_matches_gold(self, shot2, _shot_gold("shot2"))
        scallion = next(it for it in shot2["items"] if it["品名"] == "小葱")
        self.assertAlmostEqual(scallion["数量"], 8.0, places=2)
        self.assertAlmostEqual(scallion["金额"], 71.31, places=2)
        self.assertFalse(scallion["实收缺失"])
        self.assertEqual([it["品名"] for it in shot2["items"]].count("香菜"), 1)
        coriander = next(it for it in shot2["items"] if it["品名"] == "香菜")
        self.assertAlmostEqual(coriander["数量"], 3.04, places=2)
        self.assertAlmostEqual(coriander["金额"], 19.64, places=2)
        herb = next(it for it in shot2["items"] if it["品名"] == "鱼腥草")
        self.assertAlmostEqual(herb["数量"], 1.0, places=2)
        self.assertAlmostEqual(herb["金额"], 6.18, places=2)
        self.assertNotAlmostEqual(herb["数量"], 4.04, places=2)
        self.assertNotAlmostEqual(herb["金额"], 25.82, places=2)

    def test_shot1_fixture_keeps_coriander_when_price_line_moves(self):
        """同一份 Shot1 原文里，香菜的标价和订行被挪到下一件货之后，整行仍是 2/13.70。

        品名若被下一条盖掉，13.70 会写进鸡蛋或直接消失，明细和变成 244.53。
        """
        text = _read_text(SHOT_OCR["shot1"])
        moved = text.replace("￥13.70/2.00斤\n订 ￥13.70/2.00斤\n", "", 1)
        parsed = core.parse_meicai_text(moved)
        coriander = next(it for it in parsed["items"] if it["品名"] == "香菜")
        egg = next(it for it in parsed["items"] if it["品名"] == "鲜鸡蛋 白壳")
        self.assertAlmostEqual(coriander["数量"], 2.0, places=2)
        self.assertAlmostEqual(coriander["金额"], 13.70, places=2)
        self.assertFalse(coriander["实收缺失"])
        self.assertAlmostEqual(egg["数量"], 2.0, places=2)
        self.assertAlmostEqual(egg["金额"], 41.56, places=2)
        self.assertEqual(len(parsed["items"]), 6)
        self.assertFalse(any(not str(it["品名"]).strip() for it in parsed["items"]))

    def test_shot2_fixture_does_not_add_coriander_into_herb(self):
        """同一份 Shot2 原文：广告挡在错挂的「收」前面，再多一行鱼腥草。

        香菜仍是 3.04/19.64，鱼腥草仍是 1/6.18。两行小葱仍合并成 8/71.31。
        """
        text = _read_text(SHOT_OCR["shot2"])
        text = text.replace(
            "收 ￥19.64/3.04斤\n",
            "优质，拒绝欺骗！\n收 ￥19.64/3.04斤\n",
            1,
        )
        text = text.replace(
            "促鱼腥草 斤\n",
            "促鱼腥草 斤\n￥6.18/1.00斤\n订 ￥6.18/1.00斤\n鱼腥草 斤\n",
            1,
        )
        parsed = core.parse_meicai_text(text)
        scallion = next(it for it in parsed["items"] if it["品名"] == "小葱")
        self.assertAlmostEqual(scallion["数量"], 8.0, places=2)
        self.assertAlmostEqual(scallion["金额"], 71.31, places=2)
        self.assertFalse(scallion["实收缺失"])
        coriander = [it for it in parsed["items"] if it["品名"] == "香菜"]
        self.assertEqual(len(coriander), 1)
        self.assertAlmostEqual(coriander[0]["数量"], 3.04, places=2)
        self.assertAlmostEqual(coriander[0]["金额"], 19.64, places=2)
        herb = [it for it in parsed["items"] if it["品名"] == "鱼腥草"]
        self.assertEqual(len(herb), 1)
        self.assertAlmostEqual(herb[0]["数量"], 1.0, places=2)
        self.assertAlmostEqual(herb[0]["金额"], 6.18, places=2)
        self.assertFalse(any(not str(it["品名"]).strip() for it in parsed["items"]))

    def test_live_shot1_ocr_keeps_coriander_from_html_table(self):
        """本机 Shot1 原文：正文没有香菜，香菜只在 HTML 表的「香菜优质斤 / ￥13.70/2.00斤」。

        表里其余行正文已经有实收，不能再加一遍。落地仍是 6 行，香菜 2 斤 / 13.70。
        """
        text = _read_text(LIVE_OCR["shot1"])
        self.assertIn("香菜优质斤", text)
        self.assertIn("<table>", text)
        self.assertNotIn("\n香菜", text)
        parsed = core.parse_meicai_text(text)
        coriander = [it for it in parsed["items"] if it["品名"] == "香菜"]
        self.assertEqual(len(coriander), 1)
        self.assertAlmostEqual(coriander[0]["数量"], 2.0, places=2)
        self.assertAlmostEqual(coriander[0]["金额"], 13.70, places=2)
        self.assertEqual(coriander[0]["单位"], "斤")
        self.assertFalse(coriander[0]["实收缺失"])
        self.assertEqual(len(parsed["items"]), 6)
        self.assertEqual([it["品名"] for it in parsed["items"]].count("小葱"), 1)
        self.assertEqual([it["品名"] for it in parsed["items"]].count("鲜鸡蛋 白壳"), 1)
        self.assertEqual([it["品名"] for it in parsed["items"]].count("百事可乐汽水"), 1)
        egg = next(it for it in parsed["items"] if it["品名"] == "鲜鸡蛋 白壳")
        self.assertAlmostEqual(egg["数量"], 2.0, places=2)
        self.assertAlmostEqual(egg["金额"], 41.56, places=2)
        self.assertAlmostEqual(sum(it["金额"] for it in parsed["items"]), 258.23, places=2)

    def test_live_shot2_ocr_splits_coriander_from_herb(self):
        """本机 Shot2 原文：香菜和「收 ￥19.64/3.04斤」之间夹着促鱼腥草，鱼腥草另有 6.18。

        香菜必须是 3.04/19.64，鱼腥草必须是 1/6.18，禁止加总成 4.04/25.82。
        小葱的能力边界：全文只有一行标价「小葱（香葱）普通 斤 ￥26.91/3.00斤」，
        没有「收」，也没有第二档 44.40/5。标价不是实收，不能编成 8/71.31。
        三行同名收成一行，数量和金额留空。品名停在「小葱（香葱）普通」（「普通」贴在括号后，摘不掉）。
        """
        text = _read_text(LIVE_OCR["shot2"])
        self.assertIn("香菜 优质 斤", text)
        self.assertIn("收 ￥19.64/3.04斤", text)
        self.assertIn("促鱼腥草 斤", text)
        self.assertIn("￥6.18/1.00斤", text)
        self.assertNotIn("44.40", text)
        self.assertNotIn("\n收 ￥26.91", text)
        parsed = core.parse_meicai_text(text)
        coriander = [it for it in parsed["items"] if it["品名"] == "香菜"]
        self.assertEqual(len(coriander), 1)
        self.assertAlmostEqual(coriander[0]["数量"], 3.04, places=2)
        self.assertAlmostEqual(coriander[0]["金额"], 19.64, places=2)
        self.assertFalse(coriander[0]["实收缺失"])
        herb = [it for it in parsed["items"] if it["品名"] == "鱼腥草"]
        self.assertEqual(len(herb), 1)
        self.assertAlmostEqual(herb[0]["数量"], 1.0, places=2)
        self.assertAlmostEqual(herb[0]["金额"], 6.18, places=2)
        self.assertNotAlmostEqual(herb[0]["数量"], 4.04, places=2)
        self.assertNotAlmostEqual(herb[0]["金额"], 25.82, places=2)
        scallions = [it for it in parsed["items"] if "小葱" in it["品名"]]
        self.assertEqual(len(scallions), 1)
        self.assertEqual(scallions[0]["品名"], "小葱（香葱）普通")
        self.assertIsNone(scallions[0]["数量"])
        self.assertIsNone(scallions[0]["金额"])
        self.assertTrue(scallions[0]["实收缺失"])
        self.assertNotEqual(scallions[0]["数量"], 8.0)
        self.assertNotEqual(scallions[0]["金额"], 71.31)
        egg = next(it for it in parsed["items"] if it["品名"] == "鲜鸡蛋 白壳")
        self.assertAlmostEqual(egg["数量"], 2.0, places=2)
        self.assertAlmostEqual(egg["金额"], 41.56, places=2)

    def test_live_shot3_ocr_keeps_soup_bag_smaller_tier(self):
        """本机 Shot3 原文：两档煲汤袋都有「收」，仍留较小的 1 板 / 4.80，不加总成 3/14.20。"""
        text = _read_text(LIVE_OCR["shot3"])
        parsed = core.parse_meicai_text(text)
        bags = [it for it in parsed["items"] if "煲汤袋" in it["品名"]]
        self.assertEqual(len(bags), 1)
        self.assertAlmostEqual(bags[0]["数量"], 1.0, places=2)
        self.assertAlmostEqual(bags[0]["金额"], 4.80, places=2)
        self.assertEqual(bags[0]["单位"], "板")
        self.assertNotAlmostEqual(bags[0]["数量"], 3.0, places=2)
        self.assertNotAlmostEqual(bags[0]["金额"], 14.20, places=2)
        greens = next(it for it in parsed["items"] if it["品名"] == "鸡毛菜")
        self.assertAlmostEqual(greens["数量"], 3.0, places=2)
        self.assertAlmostEqual(greens["金额"], 6.54, places=2)

    def test_soda_typo_alias_and_single_promo_line_kept(self):
        """汽水谈岁饮料换成百事可乐汽水。只有促销行时去掉「促」，不把这行丢掉。"""
        text = "\n".join([
            "美菜",
            "送达时间 2026-09-02 06:00:00",
            "订单号：2996559031 复制",
            "汽水谈岁饮料 箱",
            "订 ￥120.00/3.00箱",
            "收 ￥120.00/3.00箱",
            "促鸡毛菜 斤",
            "订 ￥32.10/5.00斤",
            "收 ￥32.10/5.00斤",
            "付款信息",
            "商品多退少补合计￥152.10",
        ])
        parsed = core.parse_meicai_text(text)
        self.assertEqual([it["品名"] for it in parsed["items"]], ["百事可乐汽水", "鸡毛菜"])
        soda, greens = parsed["items"]
        self.assertAlmostEqual(soda["数量"], 3.0, places=2)
        self.assertAlmostEqual(soda["金额"], 120.0, places=2)
        self.assertEqual(soda["单位"], "箱")
        self.assertAlmostEqual(greens["数量"], 5.0, places=2)
        self.assertAlmostEqual(greens["金额"], 32.1, places=2)
        self.assertFalse(soda["实收缺失"])
        self.assertFalse(greens["实收缺失"])


def _read_text(path):
    with open(path, encoding="utf-8") as f:
        return f.read()


class MeicaiBareRecvAndNoise(unittest.TestCase):
    """长图分条 OCR：收和金额拆行、坐标贴在文案上、标价单独一行。"""

    def test_bare_recv_merges_next_line_qty(self):
        text = "\n".join([
            "美菜",
            "送达时间 2026-09-04 07:12:48",
            "订单号：2998060810 复制",
            "促 大娃娃菜 优质 斤",
            "￥6.90/3.00斤",
            "订",
            "￥6.90/3.00斤",
            "收",
            "￥6.99/3.04斤",
            "付款信息",
            "商品多退少补合计￥6.99",
        ])
        parsed = core.parse_meicai_text(text)
        self.assertEqual(len(parsed["items"]), 1)
        it = parsed["items"][0]
        self.assertEqual(it["品名"], "大娃娃菜")
        self.assertAlmostEqual(it["数量"], 3.04, places=2)
        self.assertAlmostEqual(it["金额"], 6.99, places=2)
        self.assertEqual(it["单位"], "斤")
        self.assertFalse(it["实收缺失"])
        self.assertNotAlmostEqual(it["数量"], 3.00, places=2)
        self.assertNotAlmostEqual(it["金额"], 6.90, places=2)

    def test_bare_recv_allows_spaces_and_fullwidth_yen(self):
        text = "\n".join([
            "订单号：2998060810 复制",
            "送达时间 2026-09-04 07:12:48",
            "美菜",
            "鸡毛菜 普通 斤",
            "  收  ",
            "¥6.54/3.00斤",
        ])
        it = core.parse_meicai_text(text)["items"][0]
        self.assertEqual(it["品名"], "鸡毛菜")
        self.assertAlmostEqual(it["数量"], 3.0, places=2)
        self.assertAlmostEqual(it["金额"], 6.54, places=2)
        self.assertFalse(it["实收缺失"])

    def test_noise_bbox_and_chrome_are_not_names(self):
        text = "\n".join([
            "美菜",
            "订单详情",
            "(56,141),(117,169)已完成",
            "(49,217),(537,236)",
            "美菜配送-已送达",
            "您的订单被拆分为1个包裹配送，点击查看更多物流信息",
            "送达时间 2026-09-04 07:12:48",
            "签收方式 信任签收 | 店门口",
            "司机姓名：苏师傅 联系司机",
            "退款详情",
            "订单号：2998060810 复制",
            "大娃娃菜 优质 斤",
            "收",
            "￥1.50/1.00斤",
            "付款信息",
            "商品多退少补合计￥1.50",
        ])
        parsed = core.parse_meicai_text(text)
        names = [it["品名"] for it in parsed["items"]]
        self.assertEqual(names, ["大娃娃菜"])
        self.assertAlmostEqual(parsed["items"][0]["数量"], 1.0, places=2)
        blob = "\n".join(names)
        for word in ("已完成", "美菜配送", "退款详情", "送达时间", "信任签收", "查看更多物流", "司机姓名", "订单详情"):
            self.assertNotIn(word, blob)
        self.assertIsNone(re.search(r"\(\s*\d+\s*,\s*\d+\s*\)", blob))
        self.assertEqual(parsed["date"], "2026-09-04")

    def test_missing_recv_keeps_name_when_price_is_on_next_line(self):
        text = "\n".join([
            "美菜",
            "送达时间 2026-09-04 07:12:48",
            "￥6.90",
            "订单号：2998060810 复制",
            "大娃娃菜 优质 斤",
            "￥6.90/3.00斤",
            "订",
            "￥6.90/3.00斤",
            "付款信息",
            "商品多退少补合计￥6.90",
        ])
        parsed = core.parse_meicai_text(text)
        self.assertEqual(len(parsed["items"]), 1)
        it = parsed["items"][0]
        self.assertEqual(it["品名"], "大娃娃菜")
        self.assertIsNone(it["数量"])
        self.assertIsNone(it["金额"])
        self.assertTrue(it["实收缺失"])
        self.assertTrue(any("实收" in s for s in parsed["suspects"]))
        self.assertNotEqual(it["数量"], 3.0)
        self.assertNotEqual(it["金额"], 6.9)

    def test_multi_segment_pending_kept_when_next_name_arrives(self):
        """连续两段「品名 + 下一行标价」都没有「收」，第三条品名不能把前两段盖掉。"""
        text = "\n".join([
            "美菜",
            "送达时间 2026-09-04 07:12:48",
            "订单号：2998060810 复制",
            "大娃娃菜 优质 斤",
            "￥6.90/3.00斤",
            "订",
            "￥6.90/3.00斤",
            "鸡毛菜 普通 斤",
            "￥6.54/3.00斤",
            "订",
            "￥6.54/3.00斤",
            "鲜鸡蛋 红壳 中码 托(30个)",
            "付款信息",
            "商品多退少补合计￥20.00",
        ])
        parsed = core.parse_meicai_text(text)
        self.assertEqual(
            [it["品名"] for it in parsed["items"]],
            ["大娃娃菜", "鸡毛菜", "鲜鸡蛋 红壳"],
        )
        for it in parsed["items"]:
            self.assertIsNone(it["数量"])
            self.assertIsNone(it["金额"])
            self.assertTrue(it["实收缺失"])
        self.assertNotEqual(parsed["items"][0]["数量"], 3.0)
        self.assertNotEqual(parsed["items"][0]["金额"], 6.9)
        self.assertNotEqual(parsed["items"][1]["数量"], 3.0)
        self.assertNotEqual(parsed["items"][1]["金额"], 6.54)
        self.assertTrue(any("实收" in s for s in parsed["suspects"]))
        self.assertEqual(parsed["paid"], 20.0)

    def test_shot3_tiled_fixture_fills_recv_and_drops_noise(self):
        text = _read_text(SHOT3_TILED_OCR)
        parsed = core.parse_meicai_text(text)
        self.assertTrue(core.is_meicai_ocr(text))
        self.assertEqual(parsed["date"], "2026-09-04")
        self.assertEqual(parsed["paid"], 87.34)
        cabbage = next(it for it in parsed["items"] if it["品名"] == "大娃娃菜")
        self.assertAlmostEqual(cabbage["数量"], 3.04, places=2)
        self.assertAlmostEqual(cabbage["金额"], 6.99, places=2)
        self.assertFalse(cabbage["实收缺失"])
        greens = next(it for it in parsed["items"] if it["品名"] == "鸡毛菜")
        self.assertAlmostEqual(greens["数量"], 3.0, places=2)
        self.assertAlmostEqual(greens["金额"], 6.54, places=2)
        egg = next(it for it in parsed["items"] if it["品名"] == "鲜鸡蛋 红壳")
        self.assertAlmostEqual(egg["数量"], 1.0, places=2)
        self.assertAlmostEqual(egg["金额"], 20.08, places=2)
        self.assertEqual(egg["单位"], "托")
        claw = next(it for it in parsed["items"] if it["品名"] == "冻鸡爪 L号 大爪/长爪")
        self.assertAlmostEqual(claw["数量"], 1.0, places=2)
        self.assertAlmostEqual(claw["金额"], 28.69, places=2)
        self.assertEqual(claw["单位"], "袋")
        herb = next(it for it in parsed["items"] if it["品名"] == "鱼腥草")
        self.assertAlmostEqual(herb["数量"], 1.0, places=2)
        self.assertAlmostEqual(herb["金额"], 6.18, places=2)
        # 香菜这条「收」是 0/0，但品名在。留下给人工看，不编成订货数 15.04。
        coriander = next(it for it in parsed["items"] if it["品名"] == "香菜")
        self.assertAlmostEqual(coriander["数量"], 0.0, places=2)
        self.assertAlmostEqual(coriander["金额"], 0.0, places=2)
        self.assertFalse(coriander["实收缺失"])
        beans = [it for it in parsed["items"] if it["品名"] == "优质黄豆" and it["金额"] == 4.66]
        self.assertEqual(len(beans), 1)
        self.assertEqual(beans[0]["单位"], "袋")
        names = [it["品名"] for it in parsed["items"]]
        blob = "\n".join(names)
        for word in ("已完成", "美菜配送", "退款详情", "送达时间", "信任签收", "查看更多物流", "司机姓名"):
            self.assertNotIn(word, blob)
        self.assertIsNone(re.search(r"\(\s*\d+\s*,\s*\d+\s*\)", blob))
        # 订货价 6.90/3.00 不是大娃娃菜的实收
        self.assertFalse(any(
            it["品名"] == "大娃娃菜" and it["金额"] == 6.9 and it["数量"] == 3.0
            for it in parsed["items"]
        ))

    def test_shot3_full_image_does_not_invent_recv(self):
        text = _read_text(SHOT3_OCR)
        parsed = core.parse_meicai_text(text)
        self.assertTrue(core.is_meicai_ocr(text))
        self.assertEqual(parsed["date"], "2026-09-04")
        self.assertEqual(parsed["paid"], 87.34)
        self.assertGreaterEqual(len(parsed["items"]), 8)
        self.assertTrue(any("大娃娃菜" in it["品名"] for it in parsed["items"]))
        self.assertNotIn("共9类", [it["品名"] for it in parsed["items"]])
        for it in parsed["items"]:
            self.assertTrue(it["品名"])
            self.assertIsNone(it["数量"])
            self.assertIsNone(it["金额"])
            self.assertTrue(it["实收缺失"])

    def test_quality_prefers_tiled_shot3_and_more_recv(self):
        tiled = _read_text(SHOT3_TILED_OCR)
        full = _read_text(SHOT3_OCR)
        self.assertGreater(core.meicai_ocr_quality(tiled), core.meicai_ocr_quality(full))
        self.assertIs(core.choose_meicai_ocr_text([tiled, full]), tiled)
        # 同分留前面的；整图实收更多时选整图。选的是原文，金额仍由解析读「收」。
        self.assertIs(core.choose_meicai_ocr_text([full, full]), full)
        weak = "\n".join([
            "美菜配送",
            "送达时间 2026-09-04 07:12:48",
            "订单号：2998060810 复制",
            "(1,2),(3,4)",
            "大娃娃菜 斤",
            "￥6.90/3.00斤",
            "多退少补",
        ])
        rich = "\n".join([
            "美菜",
            "送达时间 2026-09-04 07:12:48",
            "订单号：2998060810 复制",
            "大娃娃菜 斤 ￥6.90/3.00斤",
            "订 ￥6.90/3.00斤",
            "收 ￥6.99/3.04斤",
            "鸡毛菜 斤 ￥6.54/3.00斤",
            "收 ￥6.54/3.00斤",
            "商品多退少补合计￥13.53",
        ])
        self.assertGreater(core.meicai_ocr_quality(rich), core.meicai_ocr_quality(weak))
        chosen = core.choose_meicai_ocr_text([weak, rich])
        self.assertIs(chosen, rich)
        item = next(it for it in core.parse_meicai_text(chosen)["items"] if it["品名"] == "大娃娃菜")
        self.assertAlmostEqual(item["金额"], 6.99, places=2)
        self.assertAlmostEqual(item["数量"], 3.04, places=2)


class MeicaiOverlapAndAdjust(unittest.TestCase):
    """分条重叠的近重复行，以及「补 +」「退 -」角标。"""

    def test_join_drops_overlap_near_duplicate_houttuynia(self):
        """鱼腥草在相邻条里各出现一次，字不完全相同。拼完仍是 1 斤 / 6.18，不是 2 / 12.36。"""
        first = "\n".join([
            "美菜",
            "送达时间 2026-09-04 07:12:48",
            "订单号：2998060811 复制",
            "香菜 优质 斤",
            "￥15.04/2.00斤",
            "订",
            "￥15.04/2.00斤",
            "收",
            "￥15.04/2.00斤",
            "促 鱼腥草 斤",
            "￥6.18/1.00斤",
            "订",
            "￥6.18/1.00斤",
            "收",
            "￥6.18/1.00斤",
            "等额收货",
        ])
        second = "\n".join([
            "订单号：2998060811",
            "复制",
            "鱼腥草 斤",
            "￥6.18/1.00斤",
            "订 ￥6.18/1.00斤",
            "收 ¥6.18/1.00斤",
            "等额收货",
            "订单号：2998060812 复制",
            "优质黄豆 袋(500g)",
            "￥4.66/1.00袋",
            "订",
            "￥4.66/1.00袋",
            "收 ￥4.66/1.00袋",
            "付款信息",
            "商品多退少补合计￥25.88",
        ])
        # 行并不逐字相同，旧的整段相等删不掉这一份。
        self.assertNotIn("鱼腥草 斤\n￥6.18/1.00斤\n订 ￥6.18/1.00斤", first)
        parsed = core.parse_meicai_text(core.join_ocr_slices([first, second]))
        herb = [it for it in parsed["items"] if it["品名"] == "鱼腥草"]
        self.assertEqual(len(herb), 1)
        self.assertAlmostEqual(herb[0]["数量"], 1.0, places=2)
        self.assertAlmostEqual(herb[0]["金额"], 6.18, places=2)
        self.assertNotAlmostEqual(herb[0]["数量"], 2.0, places=2)
        self.assertNotAlmostEqual(herb[0]["金额"], 12.36, places=2)
        self.assertFalse(herb[0]["实收缺失"])
        coriander = next(it for it in parsed["items"] if it["品名"] == "香菜")
        self.assertAlmostEqual(coriander["数量"], 2.0, places=2)
        self.assertAlmostEqual(coriander["金额"], 15.04, places=2)
        beans = next(it for it in parsed["items"] if it["品名"] == "优质黄豆")
        self.assertAlmostEqual(beans["数量"], 1.0, places=2)
        self.assertAlmostEqual(beans["金额"], 4.66, places=2)
        self.assertEqual(beans["单位"], "袋")

    def test_join_drops_leading_duplicate_recv_line(self):
        """重叠带只多抄了一行相同的「收」，不能把 6.18 再记到下一件货上。"""
        first = "\n".join([
            "美菜",
            "订单号：2998060811 复制",
            "鱼腥草 斤",
            "收 ￥6.18/1.00斤",
        ])
        second = "\n".join([
            "(52,91),(204,170)",
            "复制",
            "收",
            "￥6.18/1.00斤",
            "等额收货",
            "优质黄豆 袋(500g)",
            "￥4.66/1.00袋",
            "订",
            "￥4.66/1.00袋",
            "收 ￥4.66/1.00袋",
        ])
        parsed = core.parse_meicai_text(core.join_ocr_slices([first, second]))
        self.assertEqual([it["品名"] for it in parsed["items"]], ["鱼腥草", "优质黄豆"])
        herb, beans = parsed["items"]
        self.assertAlmostEqual(herb["数量"], 1.0, places=2)
        self.assertAlmostEqual(herb["金额"], 6.18, places=2)
        self.assertAlmostEqual(beans["数量"], 1.0, places=2)
        self.assertAlmostEqual(beans["金额"], 4.66, places=2)
        self.assertEqual(beans["单位"], "袋")

    def test_overlap_keeps_real_second_line_for_same_name_sum(self):
        """重叠的是 3 斤那一行；5 斤是另一行。加总仍是 8 斤 / 71.31。"""
        first = "\n".join([
            "美菜",
            "送达时间 2026-09-03 06:30:00",
            "订单号：2997320402 复制",
            "小葱（香葱）普通 斤",
            "￥26.91/3.00斤",
            "订",
            "￥26.91/3.00斤",
            "收",
            "￥26.91/3.00斤",
        ])
        second = "\n".join([
            "小葱（香葱）普通 斤 ￥26.91/3.00斤",
            "订 ￥26.91/3.00斤",
            "收 ￥26.91/3.00斤",
            "小葱（香葱）普通 斤",
            "收 ￥44.40/5.00斤",
            "付款信息",
            "商品多退少补合计￥71.31",
        ])
        parsed = core.parse_meicai_text(core.join_ocr_slices([first, second]))
        self.assertEqual(len(parsed["items"]), 1)
        row = parsed["items"][0]
        self.assertEqual(row["品名"], "小葱（香葱）普通")
        self.assertEqual(row["单位"], "斤")
        self.assertAlmostEqual(row["数量"], 8.0, places=2)
        self.assertAlmostEqual(row["金额"], 71.31, places=2)
        self.assertAlmostEqual(row["单价"], 8.91, places=2)
        self.assertFalse(row["实收缺失"])

    def test_same_receipt_inside_one_text_still_sums(self):
        """没有分条边界时，两行相同的鱼腥草仍按品名+单位加总。"""
        text = "\n".join([
            "美菜",
            "送达时间 2026-09-04 07:12:48",
            "订单号：2998060811 复制",
            "鱼腥草 斤",
            "收 ￥6.18/1.00斤",
            "鱼腥草 斤",
            "收 ￥6.18/1.00斤",
            "付款信息",
            "商品多退少补合计￥12.36",
        ])
        parsed = core.parse_meicai_text(text)
        self.assertEqual(len(parsed["items"]), 1)
        row = parsed["items"][0]
        self.assertEqual(row["品名"], "鱼腥草")
        self.assertAlmostEqual(row["数量"], 2.0, places=2)
        self.assertAlmostEqual(row["金额"], 12.36, places=2)

    def test_same_row_later_in_next_slice_still_sums(self):
        """下一片开头不是这一行时，后面再次出现的相同实收仍加总。"""
        first = "\n".join([
            "美菜",
            "订单号：2998060811 复制",
            "鱼腥草 斤",
            "收 ￥6.18/1.00斤",
            "香菜 优质 斤",
            "收 ￥15.04/2.00斤",
        ])
        second = "\n".join([
            "订单号：2998060812 复制",
            "优质黄豆 袋(500g)",
            "收 ￥4.66/1.00袋",
            "鱼腥草 斤",
            "收 ￥6.18/1.00斤",
        ])
        parsed = core.parse_meicai_text(core.join_ocr_slices([first, second]))
        herb = [it for it in parsed["items"] if it["品名"] == "鱼腥草"]
        self.assertEqual(len(herb), 1)
        self.assertAlmostEqual(herb[0]["数量"], 2.0, places=2)
        self.assertAlmostEqual(herb[0]["金额"], 12.36, places=2)
        beans = next(it for it in parsed["items"] if it["品名"] == "优质黄豆")
        self.assertAlmostEqual(beans["金额"], 4.66, places=2)

    def test_handwriting_join_does_not_drop_rows(self):
        first = "<tr><td>1</td><td>小葱</td><td>斤</td><td>3</td><td>8.97</td><td>26.91</td><td></td></tr>"
        second = "<tr><td>2</td><td>小葱</td><td>斤</td><td>5</td><td>8.88</td><td>44.40</td><td></td></tr>"
        joined = core.join_ocr_slices([first, second])
        self.assertIn("26.91", joined)
        self.assertIn("44.40", joined)
        items = core.parse_line_items(core.parse_table("<table>" + joined + "</table>"))
        self.assertEqual([it["品名"] for it in items], ["小葱", "小葱"])
        self.assertAlmostEqual(items[0]["数量"], 3.0, places=2)
        self.assertAlmostEqual(items[1]["数量"], 5.0, places=2)

    def test_adjust_badges_are_not_product_lines(self):
        """「补 +」「退 -」不是货。差额不写进数量或金额。"""
        text = "\n".join([
            "美菜",
            "送达时间 2026-09-04 07:12:48",
            "订单号：2998060810 复制",
            "促 大娃娃菜 优质 斤",
            "￥6.90/3.00斤",
            "订",
            "￥6.90/3.00斤",
            "收",
            "￥6.99/3.04斤",
            "补 +",
            "退 -",
            "补 +￥0.09/0.04斤",
            "退 -￥15.04/-2.00斤",
            "补 + ￥0.09/0.04斤",
            "退 - ￥15.04/2.00斤",
            "补",
            "+￥0.09/0.04斤",
            "退",
            "-￥15.04/-2.00斤",
            "香菜 优质 斤",
            "收",
            "￥15.04/2.00斤",
            "付款信息",
            "商品多退少补合计￥22.03",
        ])
        parsed = core.parse_meicai_text(text)
        self.assertEqual([it["品名"] for it in parsed["items"]], ["大娃娃菜", "香菜"])
        cabbage, coriander = parsed["items"]
        self.assertAlmostEqual(cabbage["数量"], 3.04, places=2)
        self.assertAlmostEqual(cabbage["金额"], 6.99, places=2)
        self.assertFalse(cabbage["实收缺失"])
        self.assertAlmostEqual(coriander["数量"], 2.0, places=2)
        self.assertAlmostEqual(coriander["金额"], 15.04, places=2)
        self.assertFalse(coriander["实收缺失"])
        for it in parsed["items"]:
            self.assertNotEqual(it["数量"], 0.04)
            self.assertNotEqual(it["金额"], 0.09)
            self.assertNotIn("补", it["品名"])
            self.assertNotIn("退", it["品名"])
        self.assertFalse(parsed["suspects"])
        self.assertEqual(parsed["paid"], 22.03)

    def test_orphan_blank_zero_recv_dropped(self):
        """游离「收 ￥0.00」没有品名。空品名且数量/金额是 0 或缺的，不进明细。

        有品名但没有「收」的仍标实收缺失。有品名且实收就是 0 的留下，不改成订货数。
        """
        text = "\n".join([
            "美菜",
            "送达时间 2026-09-04 07:12:48",
            "订单号：2998060811 复制",
            "鱼腥草 斤",
            "收 ￥6.18/1.00斤",
            "收 ￥0.00/0.00斤",
            "收",
            "￥0.00/0.00",
            "收 ￥0.00",
            "收 ¥0.00/0.00斤",
            "黄瓜 二级 斤",
            "￥11.84/5.08斤",
            "香菜 优质 斤",
            "收 ￥0.00/0.00斤",
            "付款信息",
            "商品多退少补合计￥6.18",
        ])
        parsed = core.parse_meicai_text(text)
        self.assertEqual([it["品名"] for it in parsed["items"]], ["鱼腥草", "黄瓜 二级", "香菜"])
        herb, cucumber, coriander = parsed["items"]
        self.assertAlmostEqual(herb["数量"], 1.0, places=2)
        self.assertAlmostEqual(herb["金额"], 6.18, places=2)
        self.assertFalse(herb["实收缺失"])
        self.assertIsNone(cucumber["数量"])
        self.assertIsNone(cucumber["金额"])
        self.assertTrue(cucumber["实收缺失"])
        self.assertAlmostEqual(coriander["数量"], 0.0, places=2)
        self.assertAlmostEqual(coriander["金额"], 0.0, places=2)
        self.assertFalse(coriander["实收缺失"])
        self.assertNotAlmostEqual(coriander["金额"], 11.84, places=2)
        for it in parsed["items"]:
            self.assertTrue(str(it["品名"]).strip())
        self.assertEqual(parsed["suspects"], [core.meicai_missing_recv_summary(1)])

    def test_join_orphan_zero_recv_does_not_change_houttuynia(self):
        """交界上多出来的「收 ￥0.00」对不上上一笔实收，拼接留着，解析时丢掉。"""
        first = "\n".join([
            "美菜",
            "送达时间 2026-09-04 07:12:48",
            "订单号：2998060811 复制",
            "鱼腥草 斤",
            "收 ￥6.18/1.00斤",
        ])
        second = "\n".join([
            "收 ￥0.00/0.00斤",
            "收",
            "￥0.00",
            "优质黄豆 袋(500g)",
            "￥4.66/1.00袋",
            "订",
            "￥4.66/1.00袋",
            "收 ￥4.66/1.00袋",
            "付款信息",
            "商品多退少补合计￥10.84",
        ])
        joined = core.join_ocr_slices([first, second])
        self.assertIn("收 ￥0.00/0.00斤", joined)
        parsed = core.parse_meicai_text(joined)
        self.assertEqual([it["品名"] for it in parsed["items"]], ["鱼腥草", "优质黄豆"])
        herb, beans = parsed["items"]
        self.assertAlmostEqual(herb["数量"], 1.0, places=2)
        self.assertAlmostEqual(herb["金额"], 6.18, places=2)
        self.assertFalse(herb["实收缺失"])
        self.assertAlmostEqual(beans["数量"], 1.0, places=2)
        self.assertAlmostEqual(beans["金额"], 4.66, places=2)
        self.assertEqual(beans["单位"], "袋")
        self.assertFalse(beans["实收缺失"])
        self.assertFalse(parsed["suspects"])

    def test_blank_name_nonzero_recv_kept_without_inventing_name(self):
        """空品名但金额不是 0：留下这笔实收，不编品名，也不写到后面的货上。"""
        text = "\n".join([
            "美菜",
            "送达时间 2026-09-04 07:12:48",
            "订单号：2998060811 复制",
            "收 ￥6.18/1.00斤",
            "鱼腥草 斤",
            "付款信息",
            "商品多退少补合计￥6.18",
        ])
        parsed = core.parse_meicai_text(text)
        self.assertEqual(len(parsed["items"]), 2)
        blank, herb = parsed["items"]
        self.assertEqual(blank["品名"], "")
        self.assertAlmostEqual(blank["数量"], 1.0, places=2)
        self.assertAlmostEqual(blank["金额"], 6.18, places=2)
        self.assertFalse(blank["实收缺失"])
        self.assertEqual(herb["品名"], "鱼腥草")
        self.assertIsNone(herb["数量"])
        self.assertIsNone(herb["金额"])
        self.assertTrue(herb["实收缺失"])

    def test_handwriting_keeps_literal_adjust_text(self):
        html = "\n".join([
            "<table><tr><td>编号</td><td>商品名称</td><td>单位</td><td>数量</td><td>单价</td><td>金额</td><td>备注</td></tr>",
            "<tr><td>1</td><td>补 +</td><td>斤</td><td>1</td><td>2</td><td>2</td><td></td></tr>",
            "</table>",
        ])
        self.assertFalse(core.is_meicai_ocr(html))
        items = core.parse_line_items(core.parse_table(html))
        self.assertEqual(items[0]["品名"], "补 +")
        self.assertAlmostEqual(items[0]["数量"], 1.0, places=2)
        self.assertAlmostEqual(items[0]["金额"], 2.0, places=2)

    def test_dense_slice_paren_variants_stay_one_item(self):
        """更密的相邻条把同一件货认了好几遍，括号和尺寸写法不同。

        煲汤袋仍是 1 板，不是 3；冻鸡爪仍是 1 袋，不是 2。
        后面数量不同的香菜是另一行，加总仍按这一行的实收。
        """
        def goods(rows):
            lines = []
            for name, amt, qty, unit in rows:
                lines += [
                    name,
                    f"￥{amt}/{qty}{unit}",
                    "订",
                    f"￥{amt}/{qty}{unit}",
                    "收",
                    f"￥{amt}/{qty}{unit}",
                ]
            return lines

        first = "\n".join([
            "美菜",
            "送达时间 2026-09-04 07:12:48",
            "订单号：2998060809 复制",
            *goods([
                ("无纺布煲汤袋 10*15cm 板(12个)", "4.80", "1.00", "板"),
                ("[雪靓] 冻鸡爪 L号 大爪/长爪 袋(1Kg)", "28.69", "1.00", "袋"),
                ("大娃娃菜 优质 斤", "6.99", "3.04", "斤"),
                ("鸡毛菜 普通 斤", "6.54", "3.00", "斤"),
                ("鲜鸡蛋 红壳 中码 托(30个)", "20.08", "1.00", "托"),
            ]),
        ])
        second = "\n".join([
            "优质，拒绝欺骗！",
            *goods([
                ("无纺布煲汤袋(12个) 板", "4.80", "1.00", "板"),
                ("[雪靓]冻鸡爪L号 大爪/长爪 袋(1Kg)", "28.69", "1.00", "袋"),
                ("大娃娃菜 优质 斤", "6.99", "3.04", "斤"),
                ("鸡毛菜 普通 斤", "6.54", "3.00", "斤"),
                ("鲜鸡蛋 红壳 中码 托(30个)", "20.08", "1.00", "托"),
                ("香菜 优质 斤", "15.04", "2.00", "斤"),
            ]),
        ])
        third = "\n".join([
            "由于季节和运输原因，会有烧心、碰碰黑斑现象。优质，拒绝欺骗！",
            *goods([
                ("无纺布煲汤袋 板（12个）", "4.80", "1.00", "板"),
                ("冻鸡爪 L号 大爪/长爪 袋(1Kg)", "28.69", "1.00", "袋"),
                ("大娃娃菜 优质 斤", "6.99", "3.04", "斤"),
                ("鸡毛菜 普通 斤", "6.54", "3.00", "斤"),
                ("鲜鸡蛋 红壳 中码 托(30个)", "20.08", "1.00", "托"),
                ("香菜 优质 斤", "15.04", "2.00", "斤"),
                ("黄瓜 二级 斤", "11.84", "5.08", "斤"),
            ]),
            "付款信息",
            "商品多退少补合计￥93.00",
        ])
        parsed = core.parse_meicai_text(core.join_ocr_slices([first, second, third]))
        bags = [it for it in parsed["items"] if "煲汤" in it["品名"]]
        claws = [it for it in parsed["items"] if "鸡爪" in it["品名"]]
        self.assertEqual(len(bags), 1)
        self.assertEqual(bags[0]["单位"], "板")
        self.assertAlmostEqual(bags[0]["数量"], 1.0, places=2)
        self.assertAlmostEqual(bags[0]["金额"], 4.80, places=2)
        self.assertNotAlmostEqual(bags[0]["数量"], 3.0, places=2)
        self.assertEqual(len(claws), 1)
        self.assertEqual(claws[0]["单位"], "袋")
        self.assertAlmostEqual(claws[0]["数量"], 1.0, places=2)
        self.assertAlmostEqual(claws[0]["金额"], 28.69, places=2)
        self.assertNotAlmostEqual(claws[0]["数量"], 2.0, places=2)
        cabbage = next(it for it in parsed["items"] if it["品名"] == "大娃娃菜")
        self.assertAlmostEqual(cabbage["数量"], 3.04, places=2)
        self.assertAlmostEqual(cabbage["金额"], 6.99, places=2)
        coriander = next(it for it in parsed["items"] if it["品名"] == "香菜")
        self.assertAlmostEqual(coriander["数量"], 2.0, places=2)
        self.assertAlmostEqual(coriander["金额"], 15.04, places=2)
        cucumber = next(it for it in parsed["items"] if it["品名"] == "黄瓜 二级")
        self.assertAlmostEqual(cucumber["数量"], 5.08, places=2)
        self.assertAlmostEqual(cucumber["金额"], 11.84, places=2)
        blob = "\n".join(it["品名"] for it in parsed["items"])
        self.assertNotIn("拒绝欺骗", blob)
        self.assertNotIn("烧心", blob)

    def test_long_identical_overlap_does_not_multiply_qty(self):
        """重叠超过原先的 4 件货、行数也超过逐行比对上限时，数量仍不加倍。"""
        rows = [
            ("无纺布煲汤袋 板(12个)", "4.80", "1.00", "板"),
            ("冻鸡爪 L号 大爪/长爪 袋(1Kg)", "28.69", "1.00", "袋"),
            ("大娃娃菜 优质 斤", "6.99", "3.04", "斤"),
            ("鸡毛菜 普通 斤", "6.54", "3.00", "斤"),
            ("鲜鸡蛋 红壳 中码 托(30个)", "20.08", "1.00", "托"),
            ("香菜 优质 斤", "15.04", "2.00", "斤"),
            ("鱼腥草 斤", "6.18", "1.00", "斤"),
            ("优质黄豆 袋(500g)", "4.66", "1.00", "袋"),
        ]

        def block(items, header):
            lines = []
            if header:
                lines += ["美菜", "送达时间 2026-09-04 07:12:48", "订单号：2998060810 复制"]
            for name, amt, qty, unit in items:
                lines += [name, f"订 ￥{amt}/{qty}{unit}", f"收 ￥{amt}/{qty}{unit}", "等额收货"]
            return "\n".join(lines)

        nxt_rows = list(rows)
        nxt_rows.append(("黄瓜 二级 斤", "11.84", "5.08", "斤"))
        parsed = core.parse_meicai_text(core.join_ocr_slices([
            block(rows, True),
            block(nxt_rows, False) + "\n付款信息\n商品多退少补合计￥100",
        ]))
        expect = {
            "无纺布煲汤袋": (1.0, 4.80),
            "冻鸡爪 L号 大爪/长爪": (1.0, 28.69),
            "大娃娃菜": (3.04, 6.99),
            "鸡毛菜": (3.0, 6.54),
            "鲜鸡蛋 红壳": (1.0, 20.08),
            "香菜": (2.0, 15.04),
            "鱼腥草": (1.0, 6.18),
            "优质黄豆": (1.0, 4.66),
            "黄瓜 二级": (5.08, 11.84),
        }
        self.assertEqual(sorted(it["品名"] for it in parsed["items"]), sorted(expect))
        for it in parsed["items"]:
            qty, amt = expect[it["品名"]]
            self.assertAlmostEqual(it["数量"], qty, places=2)
            self.assertAlmostEqual(it["金额"], amt, places=2)
            self.assertFalse(it["实收缺失"])

    def test_slogan_ad_lines_do_not_become_items(self):
        """「优质，拒绝欺骗！」以及同类宣传句不进明细，也不粘进品名。

        夹在订和「收」之间、且「收」和订行一致时，实收仍记在这件货上。优质黄豆不是广告。
        """
        text = "\n".join([
            "美菜",
            "送达时间 2026-09-04 07:12:48",
            "订单号：2998060809 复制",
            "无纺布煲汤袋 10*15cm 板(12个)",
            "订 ￥4.80/1.00板",
            "优质，拒绝欺骗！",
            "由于季节和运输原因，会有烧心、碰碰黑斑现象。优质，拒绝欺骗！",
            "收",
            "￥4.80/1.00板",
            "拒绝双氧水漂白",
            "国产屠宰厂速冻锁鲜",
            "[雪靓] 冻鸡爪 L号 大爪/长爪 袋(1Kg)",
            "收 ￥28.69/1.00袋",
            "香菜 优质 斤",
            "收 ￥0.00/0.00斤",
            "茎干均匀，真优质 拒绝欺骗",
            "优质黄豆 袋(500g)",
            "收 ￥4.66/1.00袋",
            "真材实料，拒绝以次充好",
            "付款信息",
            "商品多退少补合计￥38.15",
        ])
        parsed = core.parse_meicai_text(text)
        names = [it["品名"] for it in parsed["items"]]
        self.assertEqual(names, ["无纺布煲汤袋 10*15cm", "冻鸡爪 L号 大爪/长爪", "香菜", "优质黄豆"])
        bag, claw, coriander, beans = parsed["items"]
        self.assertAlmostEqual(bag["数量"], 1.0, places=2)
        self.assertAlmostEqual(bag["金额"], 4.80, places=2)
        self.assertEqual(bag["单位"], "板")
        self.assertFalse(bag["实收缺失"])
        self.assertAlmostEqual(claw["数量"], 1.0, places=2)
        self.assertAlmostEqual(claw["金额"], 28.69, places=2)
        self.assertAlmostEqual(coriander["数量"], 0.0, places=2)
        self.assertAlmostEqual(coriander["金额"], 0.0, places=2)
        self.assertFalse(coriander["实收缺失"])
        self.assertAlmostEqual(beans["数量"], 1.0, places=2)
        self.assertAlmostEqual(beans["金额"], 4.66, places=2)
        self.assertEqual(beans["单位"], "袋")
        blob = "\n".join(names)
        for word in ("拒绝欺骗", "拒绝双氧", "烧心", "茎干均匀", "速冻锁鲜", "以次充好"):
            self.assertNotIn(word, blob)
        self.assertFalse(parsed["suspects"])

    def test_short_name_is_not_swallowed_as_paren_variant(self):
        """只差一个单位字的短品名不是重叠变体。鱼腥草和鱼腥草包都留下。"""
        first = "\n".join([
            "美菜",
            "订单号：2998060811 复制",
            "鱼腥草 斤",
            "收 ￥6.18/1.00斤",
        ])
        second = "\n".join([
            "鱼腥草包 斤",
            "收 ￥6.18/1.00斤",
            "优质黄豆 袋(500g)",
            "收 ￥4.66/1.00袋",
        ])
        parsed = core.parse_meicai_text(core.join_ocr_slices([first, second]))
        self.assertEqual([it["品名"] for it in parsed["items"]], ["鱼腥草", "鱼腥草包", "优质黄豆"])
        herb, packed, beans = parsed["items"]
        self.assertAlmostEqual(herb["数量"], 1.0, places=2)
        self.assertAlmostEqual(herb["金额"], 6.18, places=2)
        self.assertAlmostEqual(packed["数量"], 1.0, places=2)
        self.assertAlmostEqual(packed["金额"], 6.18, places=2)
        self.assertAlmostEqual(beans["数量"], 1.0, places=2)

    def test_handwriting_keeps_slogan_text(self):
        html = "\n".join([
            "<table><tr><td>编号</td><td>商品名称</td><td>单位</td><td>数量</td><td>单价</td><td>金额</td><td>备注</td></tr>",
            "<tr><td>1</td><td>优质，拒绝欺骗！</td><td>斤</td><td>1</td><td>2</td><td>2</td><td></td></tr>",
            "</table>",
        ])
        self.assertFalse(core.is_meicai_ocr(html))
        items = core.parse_line_items(core.parse_table(html))
        self.assertEqual(items[0]["品名"], "优质，拒绝欺骗！")
        self.assertAlmostEqual(items[0]["数量"], 1.0, places=2)
        self.assertAlmostEqual(items[0]["金额"], 2.0, places=2)

    def test_same_strip_bundle_tier_keeps_smaller_recv(self):
        """同一条里煲汤袋有两档实收：2 板 9.40 和 1 板 4.80。

        数量金额不同，重叠去重对不上。两档单价也不一样，不是两笔同价货。
        留下 1 板 / 4.80，不加总成 3 板 / 14.20。先后顺序不影响留哪一档。
        """
        def block(first, second):
            rows = []
            for amt, qty in (first, second):
                rows += [
                    "无纺布煲汤袋 10*15cm 板(12个)",
                    f"￥{amt}/{qty}板",
                    "订",
                    f"￥{amt}/{qty}板",
                    "收",
                    f"￥{amt}/{qty}板",
                ]
            return "\n".join([
                "美菜",
                "送达时间 2026-09-04 07:12:48",
                "订单号：2998060809 复制",
                *rows,
                "付款信息",
                "商品多退少补合计￥4.80",
            ])

        for text in (
            block(("9.40", "2.00"), ("4.80", "1.00")),
            block(("4.80", "1.00"), ("9.40", "2.00")),
        ):
            parsed = core.parse_meicai_text(text)
            bags = [it for it in parsed["items"] if "煲汤" in it["品名"]]
            self.assertEqual(len(bags), 1)
            self.assertAlmostEqual(bags[0]["数量"], 1.0, places=2)
            self.assertAlmostEqual(bags[0]["金额"], 4.80, places=2)
            self.assertEqual(bags[0]["单位"], "板")
            self.assertNotAlmostEqual(bags[0]["数量"], 3.0, places=2)
            self.assertNotAlmostEqual(bags[0]["金额"], 14.2, places=2)
            self.assertFalse(bags[0]["实收缺失"])

    def test_join_bundle_tiers_keep_one_recv(self):
        """两片各有一档煲汤袋，数量不同，拼完仍是 1 板 / 4.80。"""
        first = "\n".join([
            "美菜",
            "送达时间 2026-09-04 07:12:48",
            "订单号：2998060809 复制",
            "无纺布煲汤袋 10*15cm 板(12个)",
            "收 ￥4.80/1.00板",
        ])
        second = "\n".join([
            "无纺布煲汤袋 10*15cm 板(12个)",
            "收 ￥9.40/2.00板",
            "冻鸡爪 L号 大爪/长爪 袋(1Kg)",
            "收 ￥28.69/1.00袋",
        ])
        parsed = core.parse_meicai_text(core.join_ocr_slices([first, second]))
        bags = [it for it in parsed["items"] if "煲汤" in it["品名"]]
        self.assertEqual(len(bags), 1)
        self.assertAlmostEqual(bags[0]["数量"], 1.0, places=2)
        self.assertAlmostEqual(bags[0]["金额"], 4.80, places=2)
        claw = next(it for it in parsed["items"] if "鸡爪" in it["品名"])
        self.assertAlmostEqual(claw["数量"], 1.0, places=2)
        self.assertAlmostEqual(claw["金额"], 28.69, places=2)

    def test_same_unit_price_double_qty_still_sums(self):
        """1 板 4.80 和 2 板 9.60 单价相同，是两笔货，加总成 3 板 / 14.40。"""
        text = "\n".join([
            "美菜",
            "送达时间 2026-09-04 07:12:48",
            "订单号：2998060809 复制",
            "无纺布煲汤袋 10*15cm 板(12个)",
            "收 ￥4.80/1.00板",
            "无纺布煲汤袋 10*15cm 板(12个)",
            "收 ￥9.60/2.00板",
            "付款信息",
            "商品多退少补合计￥14.40",
        ])
        parsed = core.parse_meicai_text(text)
        bags = [it for it in parsed["items"] if "煲汤" in it["品名"]]
        self.assertEqual(len(bags), 1)
        self.assertAlmostEqual(bags[0]["数量"], 3.0, places=2)
        self.assertAlmostEqual(bags[0]["金额"], 14.40, places=2)

    def test_slogan_orphan_recv_does_not_attach_to_previous(self):
        """广告紧后面的无名「收」丢掉，不把 3.04/6.99 加进已经写过实收的鸡毛菜。"""
        text = "\n".join([
            "美菜",
            "送达时间 2026-09-04 07:12:48",
            "订单号：2998060810 复制",
            "鸡毛菜 普通 斤",
            "￥6.54/3.00斤",
            "订",
            "￥6.54/3.00斤",
            "收",
            "￥6.54/3.00斤",
            "优质，拒绝欺骗！",
            "收",
            "￥6.99/3.04斤",
            "大娃娃菜 优质 斤",
            "收 ￥6.99/3.04斤",
            "付款信息",
            "商品多退少补合计￥13.53",
        ])
        parsed = core.parse_meicai_text(text)
        self.assertEqual([it["品名"] for it in parsed["items"]], ["鸡毛菜", "大娃娃菜"])
        greens, cabbage = parsed["items"]
        self.assertAlmostEqual(greens["数量"], 3.0, places=2)
        self.assertAlmostEqual(greens["金额"], 6.54, places=2)
        self.assertNotAlmostEqual(greens["数量"], 6.04, places=2)
        self.assertNotAlmostEqual(greens["金额"], 13.53, places=2)
        self.assertFalse(greens["实收缺失"])
        self.assertAlmostEqual(cabbage["数量"], 3.04, places=2)
        self.assertAlmostEqual(cabbage["金额"], 6.99, places=2)
        self.assertFalse(any(not str(it["品名"]).strip() for it in parsed["items"]))

    def test_slogan_between_order_and_matching_recv_keeps_recv(self):
        """广告夹在订和收之间时，只把和订行一致的「收」记在这件货上。"""
        text = "\n".join([
            "美菜",
            "送达时间 2026-09-04 07:12:48",
            "订单号：2998060810 复制",
            "鸡毛菜 普通 斤",
            "订 ￥6.54/3.00斤",
            "优质，拒绝欺骗！",
            "收",
            "￥6.54/3.00斤",
            "付款信息",
            "商品多退少补合计￥6.54",
        ])
        parsed = core.parse_meicai_text(text)
        self.assertEqual(len(parsed["items"]), 1)
        it = parsed["items"][0]
        self.assertEqual(it["品名"], "鸡毛菜")
        self.assertAlmostEqual(it["数量"], 3.0, places=2)
        self.assertAlmostEqual(it["金额"], 6.54, places=2)
        self.assertFalse(it["实收缺失"])
        self.assertFalse(parsed["suspects"])

    def test_slogan_between_order_and_mismatched_recv_is_orphan(self):
        """订 3.00/6.54，广告后的 3.04/6.99 对不上，丢掉，不和后面的真鸡毛菜加总。"""
        text = "\n".join([
            "美菜",
            "送达时间 2026-09-04 07:12:48",
            "订单号：2998060810 复制",
            "鸡毛菜 普通 斤 ￥6.54/3.00斤",
            "订 ￥6.54/3.00斤",
            "优质，拒绝欺骗！",
            "收 ￥6.99/3.04斤",
            "补 +￥0.09/0.04斤",
            "鸡毛菜 普通 斤 ￥6.54/3.00斤",
            "订 ￥6.54/3.00斤",
            "收 ￥6.54/3.00斤",
            "等额收货",
            "付款信息",
            "商品多退少补合计￥6.54",
        ])
        parsed = core.parse_meicai_text(text)
        self.assertEqual([it["品名"] for it in parsed["items"]], ["鸡毛菜"])
        greens = parsed["items"][0]
        self.assertAlmostEqual(greens["数量"], 3.0, places=2)
        self.assertAlmostEqual(greens["金额"], 6.54, places=2)
        self.assertNotAlmostEqual(greens["数量"], 6.04, places=2)
        self.assertNotAlmostEqual(greens["金额"], 13.53, places=2)
        self.assertFalse(greens["实收缺失"])
        self.assertFalse(parsed["suspects"])

    def test_shot3_slogan_orphan_recv_does_not_stick_to_jimaocai(self):
        """Shot3：大娃娃菜实收之后，鸡毛菜订行、广告、交界孤儿收、再出真鸡毛菜。"""
        text = "\n".join([
            "美菜",
            "送达时间 2026-09-04 07:12:48",
            "订单号：2998060809 复制",
            "促 大娃娃菜 优质 斤 ￥6.90/3.00斤",
            "订 ￥6.90/3.00斤",
            "收 ￥6.99/3.04斤",
            "补 +￥0.09/0.04斤",
            "鸡毛菜 普通 斤 ￥6.54/3.00斤",
            "订 ￥6.54/3.00斤",
            "优质，拒绝欺骗！",
            "收 ￥6.99/3.04斤",
            "补 +￥0.09/0.04斤",
            "鸡毛菜 普通 斤 ￥6.54/3.00斤",
            "订 ￥6.54/3.00斤",
            "收 ￥6.54/3.00斤",
            "等额收货",
            "付款信息",
            "商品多退少补合计￥13.53",
        ])
        parsed = core.parse_meicai_text(text)
        self.assertEqual([it["品名"] for it in parsed["items"]], ["大娃娃菜", "鸡毛菜"])
        cabbage, greens = parsed["items"]
        self.assertAlmostEqual(cabbage["数量"], 3.04, places=2)
        self.assertAlmostEqual(cabbage["金额"], 6.99, places=2)
        self.assertFalse(cabbage["实收缺失"])
        self.assertAlmostEqual(greens["数量"], 3.0, places=2)
        self.assertAlmostEqual(greens["金额"], 6.54, places=2)
        self.assertNotAlmostEqual(greens["数量"], 6.04, places=2)
        self.assertNotAlmostEqual(greens["金额"], 13.53, places=2)
        self.assertFalse(greens["实收缺失"])
        self.assertNotIn("拒绝欺骗", "\n".join(it["品名"] for it in parsed["items"]))
        self.assertFalse(any(not str(it["品名"]).strip() for it in parsed["items"]))
        self.assertFalse(parsed["suspects"])

    def test_slogan_between_order_and_small_adjust_keeps_recv(self):
        """订 6.90/3.00，广告后收 6.99/3.04 是小额补差，仍记在这件货上。"""
        text = "\n".join([
            "美菜",
            "送达时间 2026-09-04 07:12:48",
            "订单号：2998060809 复制",
            "促 大娃娃菜 优质 斤 ￥6.90/3.00斤",
            "订 ￥6.90/3.00斤",
            "优质，拒绝欺骗！",
            "收 ￥6.99/3.04斤",
            "补 +￥0.09/0.04斤",
            "付款信息",
            "商品多退少补合计￥6.99",
        ])
        parsed = core.parse_meicai_text(text)
        self.assertEqual(len(parsed["items"]), 1)
        cabbage = parsed["items"][0]
        self.assertEqual(cabbage["品名"], "大娃娃菜")
        self.assertAlmostEqual(cabbage["数量"], 3.04, places=2)
        self.assertAlmostEqual(cabbage["金额"], 6.99, places=2)
        self.assertFalse(cabbage["实收缺失"])
        self.assertFalse(parsed["suspects"])

    def test_slogan_recv_copied_from_other_goods_is_orphan(self):
        """订行对得上，但这笔实收和前面另一件货相同，仍丢掉。同一品名的第二笔留下。"""
        copied = "\n".join([
            "美菜",
            "送达时间 2026-09-04 07:12:48",
            "订单号：2998060809 复制",
            "大娃娃菜 优质 斤 ￥6.99/3.04斤",
            "订 ￥6.99/3.04斤",
            "收 ￥6.99/3.04斤",
            "鸡毛菜 普通 斤 ￥6.99/3.04斤",
            "订 ￥6.99/3.04斤",
            "优质，拒绝欺骗！",
            "收 ￥6.99/3.04斤",
            "鸡毛菜 普通 斤 ￥6.54/3.00斤",
            "订 ￥6.54/3.00斤",
            "收 ￥6.54/3.00斤",
            "付款信息",
            "商品多退少补合计￥13.53",
        ])
        parsed = core.parse_meicai_text(copied)
        self.assertEqual([it["品名"] for it in parsed["items"]], ["大娃娃菜", "鸡毛菜"])
        cabbage, greens = parsed["items"]
        self.assertAlmostEqual(cabbage["数量"], 3.04, places=2)
        self.assertAlmostEqual(cabbage["金额"], 6.99, places=2)
        self.assertAlmostEqual(greens["数量"], 3.0, places=2)
        self.assertAlmostEqual(greens["金额"], 6.54, places=2)
        self.assertFalse(parsed["suspects"])

        same_name = "\n".join([
            "美菜",
            "送达时间 2026-09-04 07:12:48",
            "订单号：2998060811 复制",
            "鱼腥草 斤",
            "订 ￥6.18/1.00斤",
            "收 ￥6.18/1.00斤",
            "鱼腥草 斤",
            "订 ￥6.18/1.00斤",
            "优质，拒绝欺骗！",
            "收 ￥6.18/1.00斤",
            "付款信息",
            "商品多退少补合计￥12.36",
        ])
        parsed = core.parse_meicai_text(same_name)
        herbs = [it for it in parsed["items"] if it["品名"] == "鱼腥草"]
        self.assertEqual(len(herbs), 1)
        self.assertAlmostEqual(herbs[0]["数量"], 2.0, places=2)
        self.assertAlmostEqual(herbs[0]["金额"], 12.36, places=2)
        self.assertFalse(herbs[0]["实收缺失"])

    def test_gold_9_9_and_shot_scallion_do_not_regress(self):
        parsed = core.parse_meicai_text(_ocr_text())
        gold = _gold()
        self.assertEqual(parsed["date"], gold["delivery_date"])
        self.assertEqual(parsed["paid"], gold["paid"])
        self.assertEqual(parsed["order_ids"], gold["order_ids"])
        self.assertEqual(len(parsed["items"]), len(gold["日物料采购_复核后期望行"]))
        scallion = next(it for it in parsed["items"] if it["品名"] == "小葱")
        want = next(r for r in gold["日物料采购_复核后期望行"] if r["name"] == "小葱")
        self.assertAlmostEqual(scallion["数量"], want["qty"], places=2)
        self.assertAlmostEqual(scallion["金额"], want["amt"], places=2)
        self.assertFalse(scallion["实收缺失"])
        self.assertFalse(any("补" in it["品名"] or "退" in it["品名"] for it in parsed["items"]))

        shot1 = _shot_gold("shot1")
        shot2 = _shot_gold("shot2")
        s1 = next(r for r in shot1["日物料采购_复核后期望行"] if r["name"] == "小葱")
        s2 = next(r for r in shot2["日物料采购_复核后期望行"] if r["name"] == "小葱")
        self.assertAlmostEqual(s1["qty"], 5.0, places=2)
        self.assertAlmostEqual(s1["amt"], 26.9, places=2)
        self.assertAlmostEqual(s2["qty"], 8.0, places=2)
        self.assertAlmostEqual(s2["amt"], 71.31, places=2)
        merged = core.merge_meicai_docs([_doc_from_gold(shot1), _doc_from_gold(shot2)])
        self.assertEqual([d["date"] for d in merged], ["2026-09-02", "2026-09-03"])
        row = next(it for it in merged[1]["items"] if it["品名"] == "小葱")
        self.assertAlmostEqual(row["数量"], 8.0, places=2)
        self.assertAlmostEqual(row["金额"], 71.31, places=2)
        self.assertEqual(len(merged[0]["items"]), len(shot1["日物料采购_复核后期望行"]))
        self.assertEqual(len(merged[1]["items"]), len(shot2["日物料采购_复核后期望行"]))


if __name__ == "__main__":
    unittest.main()

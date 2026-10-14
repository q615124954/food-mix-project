# -*- coding: utf-8 -*-
"""员工餐分区落地（用户已确认的规则）。

- 员工餐判定：物料表里没有的品名，或位于票面员工分区之后；海底捞固定算员工餐。
- 员工餐行不调 map_product_code（大模型），不写日物料。
- 员工餐 sheet：D = 实付 − 日物料（明细），H = 日物料，I = 实付；没填实付按明细加总并提示「未扣优惠」。

OCR 原文来自 samples/gold/ocr/（真实 HunyuanOCR 返回）。8/12 员工区被阴影遮住没认出来，
测试里按「人工补录」示例补 3 行（猪肝 18、笋瓜 14、五花肉 21，用户尚未核对，只用于测试）。
"""
import copy
import json
import os
import sys
import tempfile
import unittest

from openpyxl import load_workbook

BASE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.join(BASE, "webapp"))
import app as webapp
import name_match
import pipeline_min as core

GOLD = os.path.join(BASE, "samples", "gold")
DAYS = ("11", "12", "13")
SERIAL = {d: webapp._excel_serial(f"2026-08-{d}") for d in DAYS}
STAFF_ROWS = {"11": ["海底捞", "五花肉", "去皮莴笋"], "12": ["消天老抽酱油"],
              "13": ["五花肉", "有机花菜", "油麦菜"]}
# 人工复核时的改动（金标为准）：盐被认成 10、腊肉按建议改猪腱子肉、鸡毛菜 58→5.8
HUMAN_FIX = {"10": "盐", "腊肉": "猪腱子肉"}
MANUAL_0812 = [("猪肝", 3.1, 18), ("笋瓜", 3.7, 14), ("五花肉", 1.8, 21)]
PAID = {"11": 140, "12": 122, "13": 210}


def _rows(day):
    with open(os.path.join(GOLD, "ocr", f"receipt_2026-08-{day}.ocr.txt"), encoding="utf-8") as f:
        return core.parse_table(f.read())


def _ocr_items(day):
    """解析 + 变体纠错 + 存疑判定（与 ocr_one 同顺序，不调 OCR/LLM）"""
    items, _ = core.derive_price(core.parse_line_items(_rows(day)))
    for it in items:
        raw = it["品名"]
        it["原品名"] = raw
        if raw:
            best, _c, src, _t = name_match.match(raw, webapp._legal_name_list, category="蔬菜",
                                                 units=webapp._legal_units, unit=it.get("单位"))
            if src == "variant" or src.startswith("exact"):
                it["品名"] = best
        it["纠正来源"] = ""
    suspects = webapp.compute_suspects(f"2026-08-{day}", items)
    return items, suspects


def _clear_material(src, dst):
    wb = load_workbook(src)
    for sn in ("日物料采购", "肉类采购"):
        ws = wb[sn]
        for c in range(1, ws.max_column + 1):
            if webapp._date_key(ws.cell(1, c).value) in SERIAL.values():
                for cc in (c, c + 1, c + 2):
                    for r in range(3, ws.max_row + 1):
                        ws.cell(r, cc).value = None
    ws = wb["员工餐"]
    for r in range(3, ws.max_row + 1):
        if webapp._date_key(ws.cell(r, 2).value) in SERIAL.values():
            for c in (4, 8, 9):
                ws.cell(r, c).value = None
    wb.save(dst)


class _GoldEnv(unittest.TestCase):
    material_src = os.path.join(GOLD, "material_gold.xlsx")

    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self._paths = dict(webapp._paths)
        self._code_map = copy.deepcopy(webapp._code_map)
        self._variants = name_match._VARIANTS
        self._save_code_map = webapp.save_code_map
        self._map = webapp.map_product_code
        self._chat = core.call_chat
        name_match._VARIANTS = copy.deepcopy(name_match.DEFAULT_VARIANTS)
        out = os.path.join(self.tmp.name, "out")
        os.makedirs(out)
        self.material = os.path.join(self.tmp.name, "物料明细.xlsx")
        _clear_material(self.material_src, self.material)
        webapp._paths.update(output_dir=out, material_file=self.material)
        webapp._code_map = {"肉类": {}, "蔬菜": {}}
        webapp.save_code_map = lambda: None
        webapp.load_legal_names()
        self.calls = []
        real = self._map

        def spy(raw_name, *a, **kw):
            self.calls.append(raw_name)
            return real(raw_name, *a, **kw)
        webapp.map_product_code = spy

        def no_llm(*a, **kw):
            raise AssertionError("不应调用大模型")
        core.call_chat = no_llm

    def tearDown(self):
        webapp.map_product_code = self._map
        core.call_chat = self._chat
        webapp.save_code_map = self._save_code_map
        name_match._VARIANTS = self._variants
        webapp._code_map = self._code_map
        webapp._paths.clear()
        webapp._paths.update(self._paths)
        webapp.load_legal_names()


class Classification(_GoldEnv):
    def test_seven_rows_are_staff_meal(self):
        staff_total = 0
        for d in DAYS:
            items, suspects = _ocr_items(d)
            by = {i["原品名"]: i for i in items}
            for raw in STAFF_ROWS[d]:
                it = by[raw]
                self.assertEqual(it["归属"], "员工餐", (d, raw))
                self.assertFalse(it["品名存疑"], (d, raw))
                self.assertTrue(it["归属依据"])
                staff_total += 1
            for it in items:
                if it["原品名"] in STAFF_ROWS[d]:
                    continue
                if it["归属"] == "员工餐":
                    # 8/13 的「10」（盐）和「腊肉」：品名本身待改，默认员工餐但必须标出来
                    self.assertTrue(it["归属待确认"] and it["品名存疑"], (d, it["原品名"]))
                else:
                    self.assertEqual(it["归属"], "物料", (d, it["原品名"]))
        self.assertEqual(staff_total, 7)

    def test_basis_and_confirm(self):
        items, suspects = _ocr_items("11")
        by = {i["原品名"]: i for i in items}
        self.assertIn("海底捞", by["海底捞"]["归属依据"])
        self.assertFalse(by["海底捞"]["归属待确认"])
        self.assertIn("员工分区", by["五花肉"]["归属依据"])
        self.assertFalse(by["五花肉"]["归属待确认"])
        items, suspects = _ocr_items("12")
        laochou = next(i for i in items if i["原品名"] == "消天老抽酱油")
        # 物料段里不在物料表 → 默认员工餐，但要人工确认
        self.assertTrue(laochou["归属待确认"])
        self.assertIn("物料段", laochou["归属依据"])
        self.assertTrue(any("已默认归员工餐，请确认" in s and "消天老抽酱油" in s for s in suspects))

    def test_legal_name_in_staff_zone_needs_confirm(self):
        attr, why, need = webapp.classify_item("香葱", "蔬菜", in_staff_zone=True)
        self.assertEqual((attr, need), ("员工餐", True))
        self.assertEqual(webapp.classify_item("香葱", "蔬菜")[0], "物料")
        self.assertEqual(webapp.classify_item("五花肉", "肉类")[0], "物料")

    def test_staff_zone_missing_rows(self):
        self.assertEqual(core.count_staff_zone_missing(_rows("12")), 3)
        self.assertEqual(core.count_staff_zone_missing(_rows("11")), 0)
        self.assertEqual(core.count_staff_zone_missing(_rows("13")), 0)

    def test_ocr_one_prompts_staff_zone_backfill(self):
        with open(os.path.join(GOLD, "ocr", "receipt_2026-08-12.ocr.txt"), encoding="utf-8") as f:
            text = f.read()
        old = core.call_ocr
        core.call_ocr = lambda *a, **kw: text
        try:
            res = webapp.ocr_one("receipt_2026-08-12.jpg")
        finally:
            core.call_ocr = old
        self.assertTrue(any("员工区疑似漏识别 3 行" in s and "补录" in s for s in res["suspects"]))
        self.assertFalse(any(s.startswith("疑似漏识别") for s in res["suspects"]))

    def test_no_material_table_no_mass_staff(self):
        webapp._paths["material_file"] = ""
        webapp.load_legal_names()
        self.assertEqual(webapp.classify_item("青菜", "蔬菜")[0], "物料")


def _confirm_and_land(day, paid):
    """模拟人工复核后确认：按金标改错、补录 8/12 员工区，再按 confirm 的口径写落地表"""
    items, _ = _ocr_items(day)
    for it in items:
        it["品名"] = HUMAN_FIX.get(it["品名"], it["品名"])
        if it["品名"] == "鸡毛菜":
            it["数量"] = 5.8
    if day == "12":
        for name, qty, amt in MANUAL_0812:
            items.append({"品名": name, "数量": qty, "金额": amt, "员工区": True, "人工补录": True})
    items = [webapp.normalize_item(it, "蔬菜") for it in items]
    webapp.assign_attribution(items, "蔬菜")
    core.derive_price(items)
    total = core.sum_amount(items)
    webapp.append_land(f"receipt_2026-08-{day}.jpg", f"2026-08-{day}", "蔬菜", total, paid, items)
    return items


class EndToEnd(_GoldEnv):
    def _sync(self, paid):
        webapp._state = {"invoices": [{"filename": f"receipt_2026-08-{d}.jpg", "category": "蔬菜"}
                                      for d in DAYS]}
        for d in DAYS:
            _confirm_and_land(d, paid.get(d))
        return webapp.sync_to_material("unused-model")

    def _staff_sheet(self):
        wb = load_workbook(self.material, data_only=True)
        ws = wb["员工餐"]
        got = {}
        for r in range(3, ws.max_row + 1):
            key = webapp._date_key(ws.cell(r, 2).value)
            for d, sv in SERIAL.items():
                if key == sv:
                    got[d] = (ws.cell(r, 4).value, ws.cell(r, 8).value, ws.cell(r, 9).value)
        return got

    def test_gold_staff_sheet_and_daily_cells(self):
        results, _ = self._sync(PAID)
        self.assertEqual(self._staff_sheet(), {"11": (40, 100, 140), "12": (64, 58, 122),
                                               "13": (43, 167, 210)})
        # 员工餐行一次都没调 map_product_code
        staff_names = {n for v in STAFF_ROWS.values() for n in v} | {m[0] for m in MANUAL_0812}
        self.assertEqual(sorted(set(self.calls) & staff_names), [])
        # 日物料 27 格逐格对金标，其余格子（含生抽）这三天都是空的
        with open(os.path.join(GOLD, "baseline_aug11_13.json"), encoding="utf-8") as f:
            gold = json.load(f)["日物料采购"]
        self.assertEqual(len(gold), 27)
        want = {(g["date"], g["name"]): (g["qty"], g["amt"]) for g in gold}
        wb = load_workbook(self.material, data_only=True)
        ws = wb["日物料采购"]
        cols = {c: d for c in range(1, ws.max_column + 1) for d, sv in SERIAL.items()
                if webapp._date_key(ws.cell(1, c).value) == sv}
        got = {}
        for r in range(3, ws.max_row + 1):
            name = ws.cell(r, 4).value
            for c, d in cols.items():
                q, a = ws.cell(r, c).value, ws.cell(r, c + 2).value
                if q is not None or a is not None:
                    got[(f"2026-08-{d}", name)] = (q, a)
        self.assertEqual(got, want)
        self.assertFalse(any(name == "生抽" for (_d, name) in got))

    def test_no_paid_falls_back_with_hint(self):
        results, _ = self._sync({})
        got = self._staff_sheet()
        # 明细加总：8/11 141−100=41，8/12 71+53−58=66，8/13 212−167=45
        self.assertEqual(got, {"11": (41, 100, 141), "12": (66, 58, 124), "13": (45, 167, 212)})
        tips = [r["理由"] for r in results if r.get("目标品名") == "员工餐" and r.get("来源") == "mapping"]
        self.assertEqual(len(tips), 3)
        self.assertTrue(all("未扣优惠" in t for t in tips))


if __name__ == "__main__":
    unittest.main()

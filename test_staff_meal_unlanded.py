# -*- coding: utf-8 -*-
"""员工餐只统计没写进物料表的金额；海底捞不进日物料采购。"""
import os
import sys
import tempfile
import unittest

from openpyxl import Workbook, load_workbook

sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), "webapp"))
import app as webapp


DATE = "2026-08-01"


def _fake_map(calls):
    """青菜/白菜/海底捞都能精确归到日物料行。野草归并不上，用来模拟未落地。"""
    codes = {
        "青菜": "V-QING",
        "白菜": "V-BAI",
        "海底捞": "V-HDL",
    }

    def fake(raw_name, category, name_code, model, qty=None, amt=None):
        calls.append((raw_name, category))
        code = codes.get(raw_name, "")
        if not code or code not in set(name_code.values()):
            return {"物料代码": "", "目标品名": raw_name, "置信度": 0,
                    "理由": "无法归并", "来源": "error"}
        return {"物料代码": code, "目标品名": raw_name, "置信度": 1.0,
                "理由": "test", "来源": "mapping"}

    return fake


def _build_material(path):
    """日物料里预放青菜/白菜/海底捞，数量和金额先写成 99/999，用来确认没被误写。

    员工餐 sheet 同一天的金额也先写成 999，避免「在旧值上加」碰巧通过。
    """
    serial = webapp._excel_serial(DATE)
    wb = Workbook()
    veg = wb.active
    veg.title = "日物料采购"
    veg.cell(1, 2, "代码")
    veg.cell(1, 4, "品名")
    veg.cell(1, 5, "单位")
    veg.cell(1, 6, serial)
    products = [("V-QING", "青菜"), ("V-BAI", "白菜"), ("V-HDL", "海底捞")]
    for i, (code, name) in enumerate(products):
        row = 3 + i
        veg.cell(row, 2, code)
        veg.cell(row, 4, name)
        veg.cell(row, 5, "斤")
        veg.cell(row, 6, 99)
        veg.cell(row, 8, 999)

    staff = wb.create_sheet("员工餐")
    # 表头：员工餐金额在前，日物料金额在后，总计在最后
    headers = ["序号", "日期", "类别", "金额", "序号", "日期", "类别", "金额", "总计"]
    for col, text in enumerate(headers, start=1):
        staff.cell(2, col, text)
    staff.cell(3, 2, serial)
    staff.cell(3, 4, 999)   # 员工餐·金额
    staff.cell(3, 6, serial)
    staff.cell(3, 8, 999)   # 日物料·金额
    staff.cell(3, 9, 999)   # 总计
    wb.save(path)


def _build_land(path, data_rows):
    wb = Workbook()
    ws = wb.active
    ws.title = "已落地"
    ws.append(["日期", "类别", "品名", "数量", "单价", "金额", "来源图片", "合计", "实付合计"])
    for row in data_rows:
        ws.append(list(row))
    wb.save(path)


def _product_cell(ws, name):
    for r in range(3, ws.max_row + 1):
        if ws.cell(r, 4).value == name:
            return ws.cell(r, 6).value, ws.cell(r, 8).value
    return None, None


class StaffMealUnlandedTest(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self._paths = dict(webapp._paths)
        self._code_map = dict(webapp._code_map)
        out = os.path.join(self.tmp.name, "out")
        os.makedirs(out)
        webapp._paths["output_dir"] = out
        webapp._paths["material_file"] = os.path.join(self.tmp.name, "物料明细.xlsx")
        webapp._code_map = {"肉类": {}, "蔬菜": {}}
        self.calls = []
        self._map = webapp.map_product_code
        webapp.map_product_code = _fake_map(self.calls)

    def tearDown(self):
        webapp.map_product_code = self._map
        webapp._paths.clear()
        webapp._paths.update(self._paths)
        webapp._code_map.clear()
        webapp._code_map.update(self._code_map)

    def _sync(self, rows):
        _build_material(webapp._paths["material_file"])
        _build_land(os.path.join(webapp._paths["output_dir"], webapp.LAND_FILENAME), rows)
        return webapp.sync_to_material("unused-model")

    def _read(self):
        wb = load_workbook(webapp._paths["material_file"])
        veg = wb["日物料采购"]
        staff = wb["员工餐"]
        return {
            "青菜": _product_cell(veg, "青菜"),
            "白菜": _product_cell(veg, "白菜"),
            "海底捞": _product_cell(veg, "海底捞"),
            "员工餐": staff.cell(3, 4).value,
            "日物料": staff.cell(3, 8).value,
            "总计": staff.cell(3, 9).value,
        }

    def test_unlanded_rows_do_not_overstate_staff_meal(self):
        """员工餐 = 蔬菜单实付 − 本单已落日物料（用户确认的口径：优惠、差额都归员工餐）。

        ok.jpg 实付 120，只有青菜 50 写进物料表 → 员工餐 70（含没归并上的野草 40 和差额 30）。
        缺类别的白菜 15 不写物料表，计员工餐；空日期的 20 元不能记到 8 月 1 日，否则虚高。
        """
        results, _backup = self._sync([
            (DATE, "蔬菜", "青菜", 2, 25, 50, "ok.jpg", 90, 120),
            (DATE, "蔬菜", "野草", 1, 40, 40, "ok.jpg", 90, 120),
            (DATE, None, "白菜", 4, 3.75, 15, "nocat.jpg", 15, 15),
            ("", "蔬菜", "青菜", 8, 2.5, 20, "blank.jpg", 20, 20),
        ])

        got = self._read()
        # 只有成功写入的青菜进日物料；预填的 99/999 被整天合计覆盖成 2/50
        self.assertEqual(got["青菜"], (2, 50))
        self.assertEqual(got["白菜"], (99, 999))
        self.assertEqual(got["日物料"], 50)
        # ok.jpg：120 − 50 = 70；缺类别白菜 15。空日期的 20 不算进来（否则是 105）
        self.assertEqual(got["员工餐"], 85)
        self.assertEqual(got["总计"], 135)

        blank = [r for r in results if r.get("来源") == "staff_meal" and r.get("金额") == 20]
        self.assertEqual(len(blank), 1)
        self.assertEqual(blank[0]["日期"], "")
        self.assertEqual(blank[0]["落地品名"], "")
        self.assertIn("青菜", blank[0]["理由"])
        self.assertIn("不记入其他日期", blank[0]["理由"])
        self.assertNotIn(("海底捞", "蔬菜"), self.calls)

    def test_haidilao_stays_off_material_sheet_and_counts_as_staff_meal(self):
        """海底捞即使能归到日物料行，也不写数量/金额，金额进员工餐。"""
        self._sync([
            (DATE, "蔬菜", "青菜", 5, 10, 50, "veg.jpg", 90, 90),
            (DATE, "蔬菜", "海底捞", 1, 40, 40, "veg.jpg", 90, 90),
        ])
        got = self._read()
        self.assertEqual(got["青菜"], (5, 50))
        self.assertEqual(got["海底捞"], (99, 999))
        self.assertEqual(got["日物料"], 50)
        self.assertEqual(got["员工餐"], 40)
        self.assertNotIn("海底捞", [name for name, _cat in self.calls])

        # 再同步一次仍然是 50/40，海底捞行不被写上
        webapp.sync_to_material("unused-model")
        got = self._read()
        self.assertEqual(got["青菜"], (5, 50))
        self.assertEqual(got["海底捞"], (99, 999))
        self.assertEqual(got["日物料"], 50)
        self.assertEqual(got["员工餐"], 40)


if __name__ == "__main__":
    unittest.main()
